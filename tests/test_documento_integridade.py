"""
Agenda Universal PRO
Arquivo: tests/test_documento_integridade.py

Valida a integridade estrutural do documento após atualização.
"""

from pathlib import Path

from docx import Document

from core.agenda_updater import AgendaUpdater


ARQUIVO_ORIGINAL = Path("modelos") / "teste_integridade.docx"
ARQUIVO_ATUALIZADO = Path("modelos") / "teste_integridade_atualizado.docx"


def criar_documento() -> None:
    """Cria documento com parágrafos e tabela."""

    documento = Document()

    documento.add_heading("Agenda Universal PRO", level=1)

    documento.add_paragraph(
        "Audiência em 05/01/2026."
    )

    documento.add_paragraph(
        "Reunião em 10/02/2026."
    )

    tabela = documento.add_table(rows=2, cols=2)

    tabela.cell(0, 0).text = "Data"
    tabela.cell(0, 1).text = "Evento"

    tabela.cell(1, 0).text = "15/03/2027"
    tabela.cell(1, 1).text = "Prazo"

    documento.save(str(ARQUIVO_ORIGINAL))


def validar_documento() -> None:
    """Valida a estrutura do documento gerado."""

    documento = Document(str(ARQUIVO_ATUALIZADO))

    assert len(documento.paragraphs) == 3
    assert len(documento.tables) == 1

    texto = "\n".join(
        p.text for p in documento.paragraphs
    )

    assert "05/01/2027" in texto
    assert "10/02/2027" in texto

    tabela = documento.tables[0]

    assert tabela.cell(1, 0).text == "15/03/2027"
    assert tabela.cell(1, 1).text == "Prazo"


def main() -> None:
    print("\n=== TESTE DE INTEGRIDADE DO DOCUMENTO ===\n")

    ARQUIVO_ORIGINAL.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    criar_documento()

    updater = AgendaUpdater()

    quantidade = updater.atualizar_ano(
        ARQUIVO_ORIGINAL,
        2026,
        2027,
        ARQUIVO_ATUALIZADO,
    )

    print(f"Datas atualizadas: {quantidade}")

    validar_documento()

    print("Estrutura preservada.")
    print("Tabela preservada.")
    print("Teste concluído com sucesso.")


if __name__ == "__main__":
    main()