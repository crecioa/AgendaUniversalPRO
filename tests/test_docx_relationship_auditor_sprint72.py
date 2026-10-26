"""
Agenda Universal PRO
Sprint 7.2 - Testes do auditor de relacionamentos DOCX.
"""

from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from services.docx_relationship_auditor import DocxRelationshipAuditor


NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)


def criar_pacote(
    caminho: Path,
    componentes: dict[str, str | bytes],
) -> None:
    """Cria um pacote ZIP temporario para testar relacionamentos."""

    with ZipFile(caminho, "w", ZIP_DEFLATED) as pacote:
        for nome, conteudo in componentes.items():
            pacote.writestr(nome, conteudo)


def criar_relacionamento(
    destino: str,
    modo: str = "Internal",
) -> str:
    """Produz XML de relacionamento para os testes."""

    return (
        f'<Relationships xmlns="{NAMESPACE}">'
        '<Relationship '
        'Id="rId1" '
        'Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/image" '
        f'Target="{destino}" '
        f'TargetMode="{modo}"/>'
        "</Relationships>"
    )


def test_documento_word_valido(tmp_path: Path) -> None:
    caminho = tmp_path / "agenda.docx"

    documento = Document()
    documento.add_paragraph("Agenda Universal PRO 2027")
    documento.save(str(caminho))

    relatorio = DocxRelationshipAuditor.auditar(caminho)

    assert relatorio.relacionamentos_validos is True
    assert relatorio.arquivos_relacionamentos >= 1
    assert relatorio.relacionamentos_internos >= 1
    assert relatorio.relacionamentos_invalidos == ()
    assert relatorio.xml_invalidos == ()


def test_relacionamento_interno_existente(tmp_path: Path) -> None:
    caminho = tmp_path / "valido.docx"

    criar_pacote(
        caminho,
        {
            "word/document.xml": "<document/>",
            "word/media/image1.png": b"imagem",
            "word/_rels/document.xml.rels": criar_relacionamento(
                "media/image1.png"
            ),
        },
    )

    relatorio = DocxRelationshipAuditor.auditar(caminho)

    assert relatorio.relacionamentos_validos is True
    assert relatorio.relacionamentos_internos == 1
    assert relatorio.relacionamentos_invalidos == ()


def test_relacionamento_interno_ausente(tmp_path: Path) -> None:
    caminho = tmp_path / "imagem_ausente.docx"

    criar_pacote(
        caminho,
        {
            "word/document.xml": "<document/>",
            "word/_rels/document.xml.rels": criar_relacionamento(
                "media/image1.png"
            ),
        },
    )

    relatorio = DocxRelationshipAuditor.auditar(caminho)

    assert relatorio.relacionamentos_validos is False
    assert relatorio.relacionamentos_internos == 1
    assert len(relatorio.relacionamentos_invalidos) == 1
    assert "word/media/image1.png" in (
        relatorio.relacionamentos_invalidos[0]
    )


def test_relacionamento_externo(tmp_path: Path) -> None:
    caminho = tmp_path / "externo.docx"

    criar_pacote(
        caminho,
        {
            "word/document.xml": "<document/>",
            "word/_rels/document.xml.rels": criar_relacionamento(
                "https://example.com/recurso",
                modo="External",
            ),
        },
    )

    relatorio = DocxRelationshipAuditor.auditar(caminho)

    assert relatorio.relacionamentos_validos is True
    assert relatorio.relacionamentos_externos == 1
    assert relatorio.relacionamentos_internos == 0


def test_xml_relacionamento_invalido(tmp_path: Path) -> None:
    caminho = tmp_path / "xml_invalido.docx"

    criar_pacote(
        caminho,
        {
            "word/document.xml": "<document/>",
            "word/_rels/document.xml.rels": (
                "<Relationships><Relationship>"
            ),
        },
    )

    relatorio = DocxRelationshipAuditor.auditar(caminho)

    assert relatorio.relacionamentos_validos is False
    assert len(relatorio.xml_invalidos) == 1


def test_relacionamento_incompleto(tmp_path: Path) -> None:
    caminho = tmp_path / "incompleto.docx"

    xml = (
        f'<Relationships xmlns="{NAMESPACE}">'
        '<Relationship Id="rId1" Target="media/image1.png"/>'
        "</Relationships>"
    )

    criar_pacote(
        caminho,
        {
            "word/document.xml": "<document/>",
            "word/_rels/document.xml.rels": xml,
        },
    )

    relatorio = DocxRelationshipAuditor.auditar(caminho)

    assert relatorio.relacionamentos_validos is False
    assert len(relatorio.relacionamentos_invalidos) == 1


def test_documento_inexistente(tmp_path: Path) -> None:
    caminho = tmp_path / "inexistente.docx"

    with pytest.raises(FileNotFoundError):
        DocxRelationshipAuditor.auditar(caminho)


def test_auditoria_preserva_arquivo(tmp_path: Path) -> None:
    caminho = tmp_path / "preservacao.docx"

    documento = Document()
    documento.add_paragraph("Documento original")
    documento.save(str(caminho))

    conteudo_original = caminho.read_bytes()

    DocxRelationshipAuditor.auditar(caminho)

    assert caminho.read_bytes() == conteudo_original
