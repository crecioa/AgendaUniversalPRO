from docx import Document

from services.formatting_auditor import FormattingAuditor


def test_auditoria_formatacao_basica_sprint71() -> None:
    documento = Document()

    paragrafo = documento.add_paragraph()

    negrito = paragrafo.add_run("Texto em negrito")
    negrito.bold = True

    italico = paragrafo.add_run("Texto em italico")
    italico.italic = True

    sublinhado = paragrafo.add_run("Texto sublinhado")
    sublinhado.underline = True

    relatorio = FormattingAuditor().auditar(documento)

    assert relatorio.paragrafos_analisados == 1
    assert relatorio.runs_analisados == 3
    assert relatorio.runs_negrito == 1
    assert relatorio.runs_italico == 1
    assert relatorio.runs_sublinhado == 1


def test_auditoria_formatacao_em_tabela_sprint71() -> None:
    documento = Document()

    tabela = documento.add_table(rows=1, cols=1)

    run = tabela.cell(0, 0).paragraphs[0].add_run(
        "Audiencia em 10/01/2027"
    )

    run.bold = True

    relatorio = FormattingAuditor().auditar(documento)

    assert relatorio.paragrafos_analisados == 1
    assert relatorio.runs_negrito == 1


def test_auditoria_formatacao_nao_modifica_xml_sprint71() -> None:
    documento = Document()

    paragrafo = documento.add_paragraph("Agenda Universal PRO")

    paragrafo.runs[0].bold = True

    xml_antes = documento.element.xml

    FormattingAuditor().auditar(documento)

    xml_depois = documento.element.xml

    assert xml_antes == xml_depois
def test_auditoria_tabela_aninhada_sprint71() -> None:
    documento = Document()

    tabela_externa = documento.add_table(rows=1, cols=1)

    celula_externa = tabela_externa.cell(0, 0)

    run_externo = celula_externa.paragraphs[0].add_run(
        "Compromisso externo"
    )
    run_externo.bold = True

    tabela_interna = celula_externa.add_table(rows=1, cols=1)

    run_interno = tabela_interna.cell(
        0, 0
    ).paragraphs[0].add_run(
        "Compromisso interno"
    )
    run_interno.italic = True

    relatorio = FormattingAuditor().auditar(documento)

    assert relatorio.runs_negrito == 1
    assert relatorio.runs_italico == 1


def test_auditoria_celulas_mescladas_sprint71() -> None:
    documento = Document()

    tabela = documento.add_table(rows=1, cols=2)

    celula_mesclada = tabela.cell(0, 0).merge(
        tabela.cell(0, 1)
    )

    run = celula_mesclada.paragraphs[0].add_run(
        "Compromisso mesclado"
    )
    run.bold = True

    relatorio = FormattingAuditor().auditar(documento)

    assert relatorio.runs_negrito == 1


def test_auditoria_tabelas_complexas_preserva_xml_sprint71() -> None:
    documento = Document()

    tabela = documento.add_table(rows=1, cols=1)

    celula = tabela.cell(0, 0)

    celula.add_table(rows=1, cols=1)

    xml_antes = documento.element.xml

    FormattingAuditor().auditar(documento)

    assert documento.element.xml == xml_antes    