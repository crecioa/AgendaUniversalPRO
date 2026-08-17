"""
Agenda Universal PRO
Arquivo: tests/test_text_updater_ano.py

Teste da atualização automática de ano do TextUpdater.
"""

from docx import Document

from services.text_updater import TextUpdater


def main() -> None:
    documento = Document()

    documento.add_paragraph("Audiência em 05/01/2026.")
    documento.add_paragraph("Reunião em 10/02/2026.")
    documento.add_paragraph("Prazo em 15/03/2027.")
    documento.add_paragraph("Texto sem data.")

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2027,
    )

    print("\n=== TESTE DO TEXTUPDATER - ATUALIZAÇÃO DE ANO ===\n")

    for paragrafo in documento.paragraphs:
        print(paragrafo.text)

    print(f"\nSubstituições realizadas: {quantidade}")

    assert quantidade == 2
    assert documento.paragraphs[0].text == "Audiência em 05/01/2027."
    assert documento.paragraphs[1].text == "Reunião em 10/02/2027."
    assert documento.paragraphs[2].text == "Prazo em 15/03/2027."

    print("\nTeste concluído com sucesso.")


if __name__ == "__main__":
    main()