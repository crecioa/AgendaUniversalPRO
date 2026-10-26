"""
Agenda Universal PRO
Sprint 7.2 - Testes de auditoria XML em documentos DOCX.
"""

from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from services.docx_xml_auditor import DocxXMLAuditor


def criar_pacote(
    caminho: Path,
    componentes: dict[str, str | bytes],
) -> None:
    """Cria um pacote ZIP temporario para os testes."""

    with ZipFile(caminho, "w", ZIP_DEFLATED) as pacote:
        for nome, conteudo in componentes.items():
            pacote.writestr(nome, conteudo)


def test_documento_word_valido(tmp_path: Path) -> None:
    caminho = tmp_path / "agenda_valida.docx"

    documento = Document()
    documento.add_paragraph("Agenda Universal PRO 2027")
    documento.save(str(caminho))

    relatorio = DocxXMLAuditor.auditar(caminho)

    assert relatorio.xml_valido is True
    assert relatorio.componentes_xml > 0
    assert relatorio.componentes_validos == relatorio.componentes_xml
    assert relatorio.componentes_invalidos == ()


def test_xml_valido(tmp_path: Path) -> None:
    caminho = tmp_path / "xml_valido.docx"

    criar_pacote(
        caminho,
        {
            "word/document.xml": "<document><texto>2027</texto></document>",
        },
    )

    relatorio = DocxXMLAuditor.auditar(caminho)

    assert relatorio.xml_valido is True
    assert relatorio.componentes_xml == 1
    assert relatorio.componentes_validos == 1


def test_xml_invalido(tmp_path: Path) -> None:
    caminho = tmp_path / "xml_invalido.docx"

    criar_pacote(
        caminho,
        {
            "word/document.xml": "<document><texto></document>",
        },
    )

    relatorio = DocxXMLAuditor.auditar(caminho)

    assert relatorio.xml_valido is False
    assert relatorio.componentes_xml == 1
    assert relatorio.componentes_validos == 0
    assert len(relatorio.componentes_invalidos) == 1
    assert "word/document.xml" in relatorio.componentes_invalidos[0]


def test_multiplos_componentes_xml(tmp_path: Path) -> None:
    caminho = tmp_path / "multiplos_xml.docx"

    criar_pacote(
        caminho,
        {
            "word/document.xml": "<document/>",
            "word/styles.xml": "<styles/>",
            "word/settings.xml": "<settings/>",
        },
    )

    relatorio = DocxXMLAuditor.auditar(caminho)

    assert relatorio.xml_valido is True
    assert relatorio.componentes_xml == 3
    assert relatorio.componentes_validos == 3


def test_componente_rels_invalido(tmp_path: Path) -> None:
    caminho = tmp_path / "rels_invalido.docx"

    criar_pacote(
        caminho,
        {
            "word/document.xml": "<document/>",
            "word/_rels/document.xml.rels": "<Relationships>",
        },
    )

    relatorio = DocxXMLAuditor.auditar(caminho)

    assert relatorio.xml_valido is False
    assert relatorio.componentes_xml == 2
    assert relatorio.componentes_validos == 1
    assert len(relatorio.componentes_invalidos) == 1
    assert "document.xml.rels" in relatorio.componentes_invalidos[0]


def test_documento_inexistente(tmp_path: Path) -> None:
    caminho = tmp_path / "inexistente.docx"

    with pytest.raises(FileNotFoundError):
        DocxXMLAuditor.auditar(caminho)


def test_auditoria_preserva_documento(tmp_path: Path) -> None:
    caminho = tmp_path / "preservacao.docx"

    documento = Document()
    documento.add_paragraph("Documento original 2026")
    documento.save(str(caminho))

    conteudo_original = caminho.read_bytes()

    relatorio = DocxXMLAuditor.auditar(caminho)

    assert relatorio.xml_valido is True
    assert caminho.read_bytes() == conteudo_original


def test_xml_valido_e_invalido_no_mesmo_pacote(
    tmp_path: Path,
) -> None:
    caminho = tmp_path / "pacote_misto.docx"

    criar_pacote(
        caminho,
        {
            "word/document.xml": "<document/>",
            "word/styles.xml": "<styles>",
            "word/settings.xml": "<settings/>",
        },
    )

    relatorio = DocxXMLAuditor.auditar(caminho)

    assert relatorio.xml_valido is False
    assert relatorio.componentes_xml == 3
    assert relatorio.componentes_validos == 2
    assert len(relatorio.componentes_invalidos) == 1
    assert "word/styles.xml" in relatorio.componentes_invalidos[0]