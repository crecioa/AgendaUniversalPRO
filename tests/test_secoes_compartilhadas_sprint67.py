from docx import Document
from docx.enum.section import WD_SECTION_START

from services.text_updater import TextUpdater


def test_cabecalho_compartilhado_entre_secoes_sprint67() -> None:
    documento = Document()

    secao1 = documento.sections[0]
    cabecalho = secao1.header.paragraphs[0]
    cabecalho.add_run("Agenda 10/01/2026")

    secao2 = documento.add_section(WD_SECTION_START.NEW_PAGE)

    assert secao2.header.is_linked_to_previous

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 1
    assert secao1.header.paragraphs[0].text == "Agenda 10/01/2027"
    assert secao2.header.paragraphs[0].text == "Agenda 10/01/2027"
