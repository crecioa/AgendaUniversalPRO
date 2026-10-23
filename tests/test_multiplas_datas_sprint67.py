from docx import Document

from services.text_updater import TextUpdater


def test_multiplas_datas_no_mesmo_paragrafo_sprint67() -> None:
    documento = Document()
    paragrafo = documento.add_paragraph()

    titulo = paragrafo.add_run("Agenda: ")
    titulo.bold = True

    paragrafo.add_run("10/01/")
    paragrafo.add_run("2026")

    separador = paragrafo.add_run(" | ")
    separador.italic = True

    paragrafo.add_run("15/05/2026")
    paragrafo.add_run(" | 10/01/2026")

    quantidade = TextUpdater().atualizar_datas(
        documento,
        {
            "10/01/2026": "10/01/2027",
            "15/05/2026": "15/05/2027",
        },
    )

    assert quantidade == 3
    assert paragrafo.text == (
        "Agenda: 10/01/2027 | 15/05/2027 | 10/01/2027"
    )
    assert titulo.bold is True
    assert separador.italic is True
