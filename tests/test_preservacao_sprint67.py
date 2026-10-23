from pathlib import Path

from docx import Document
from core.agenda_updater import AgendaUpdater


def test_preservacao_formatacao_sprint67(tmp_path: Path) -> None:
    origem = tmp_path / "agenda_original.docx"
    destino = tmp_path / "agenda_atualizada.docx"

    documento = Document()

    paragrafo = documento.add_paragraph()
    destaque = paragrafo.add_run("Audiência: ")
    destaque.bold = True
    paragrafo.add_run("10/01/2026")

    tabela = documento.add_table(rows=1, cols=2)
    tabela.cell(0, 0).text = "Reunião"
    tabela.cell(0, 1).text = "15/05/2026"

    documento.sections[0].header.paragraphs[0].text = "Agenda Universal PRO"
    documento.sections[0].footer.paragraphs[0].text = "Documento de teste"

    documento.save(str(origem))

    quantidade = AgendaUpdater().atualizar_agenda(
        arquivo=origem,
        substituicoes={
            "10/01/2026": "10/01/2027",
            "15/05/2026": "15/05/2027",
        },
        destino=destino,
    )

    atualizado = Document(str(destino))
    original = Document(str(origem))

    assert quantidade == 2
    assert "10/01/2027" in atualizado.paragraphs[0].text
    assert atualizado.tables[0].cell(0, 1).text == "15/05/2027"

    assert "10/01/2026" in original.paragraphs[0].text

    assert atualizado.sections[0].header.paragraphs[0].text == "Agenda Universal PRO"
    assert atualizado.sections[0].footer.paragraphs[0].text == "Documento de teste"

    assert atualizado.paragraphs[0].runs[0].bold is True, (
        "A atualização eliminou a formatação em negrito."
    )
