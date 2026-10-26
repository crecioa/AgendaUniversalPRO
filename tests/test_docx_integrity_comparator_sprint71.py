from pathlib import Path
import sys
from zipfile import ZIP_DEFLATED, ZipFile
from importlib import import_module

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
DocxIntegrityComparator = import_module(
    "services.docx_integrity_comparator"
).DocxIntegrityComparator


def test_comparacao_documentos_validos_sprint71(
    tmp_path: Path,
) -> None:
    original = tmp_path / "original.docx"
    atualizado = tmp_path / "atualizado.docx"

    documento = Document()
    documento.add_paragraph("Audiencia em 10/01/2026")
    documento.save(str(original))

    documento.paragraphs[0].text = "Audiencia em 10/01/2027"
    documento.save(str(atualizado))

    relatorio = DocxIntegrityComparator.comparar(
        original,
        atualizado,
    )

    assert relatorio.original_valido is True
    assert relatorio.atualizado_valido is True
    assert relatorio.estrutura_componentes_igual is True
    assert relatorio.componentes_removidos == ()
    assert relatorio.componentes_adicionados == ()


def test_comparacao_identifica_componentes_removidos_sprint71(
    tmp_path: Path,
) -> None:
    original = tmp_path / "original.docx"
    atualizado = tmp_path / "atualizado.docx"

    with ZipFile(original, "w", ZIP_DEFLATED) as pacote:
        pacote.writestr("word/document.xml", "<document/>")
        pacote.writestr("word/media/imagem1.png", b"imagem")

    with ZipFile(atualizado, "w", ZIP_DEFLATED) as pacote:
        pacote.writestr("word/document.xml", "<document/>")

    relatorio = DocxIntegrityComparator.comparar(
        original,
        atualizado,
    )

    assert relatorio.componentes_removidos == (
        "word/media/imagem1.png",
    )

    assert relatorio.estrutura_componentes_igual is False


def test_comparacao_identifica_componentes_adicionados_sprint71(
    tmp_path: Path,
) -> None:
    original = tmp_path / "original.docx"
    atualizado = tmp_path / "atualizado.docx"

    with ZipFile(original, "w", ZIP_DEFLATED) as pacote:
        pacote.writestr("word/document.xml", "<document/>")

    with ZipFile(atualizado, "w", ZIP_DEFLATED) as pacote:
        pacote.writestr("word/document.xml", "<document/>")
        pacote.writestr("word/media/imagem2.png", b"imagem")

    relatorio = DocxIntegrityComparator.comparar(
        original,
        atualizado,
    )

    assert relatorio.componentes_adicionados == (
        "word/media/imagem2.png",
    )

    assert relatorio.estrutura_componentes_igual is False


def test_comparacao_nao_modifica_documentos_sprint71(
    tmp_path: Path,
) -> None:
    original = tmp_path / "original.docx"
    atualizado = tmp_path / "atualizado.docx"

    documento = Document()
    documento.add_paragraph("Agenda Universal PRO")

    documento.save(str(original))
    documento.save(str(atualizado))

    bytes_original = original.read_bytes()
    bytes_atualizado = atualizado.read_bytes()

    DocxIntegrityComparator.comparar(
        original,
        atualizado,
    )

    assert original.read_bytes() == bytes_original
    assert atualizado.read_bytes() == bytes_atualizado