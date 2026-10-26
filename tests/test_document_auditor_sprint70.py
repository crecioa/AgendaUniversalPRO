"""
Agenda Universal PRO
Sprint 7.0 — Testes do auditor estrutural.
"""

from docx import Document

from services.document_auditor import DocumentAuditor


def test_auditoria_documento_basico_sprint70() -> None:
    documento = Document()

    documento.add_paragraph("Agenda Universal PRO")
    documento.add_paragraph("Audiência em 10/01/2026")

    documento.add_table(rows=2, cols=2)

    auditor = DocumentAuditor()
    relatorio = auditor.auditar(documento)

    assert relatorio.paragrafos == 2
    assert relatorio.tabelas == 1
    assert relatorio.secoes == 1
    assert relatorio.hyperlinks == 0
    assert relatorio.imagens == 0
    assert relatorio.campos_especiais == 0


def test_auditoria_cabecalho_rodape_sprint70() -> None:
    documento = Document()

    secao = documento.sections[0]

    secao.header.paragraphs[0].text = "Agenda Universal PRO"
    secao.footer.paragraphs[0].text = "Página de controle"

    auditor = DocumentAuditor()
    relatorio = auditor.auditar(documento)

    assert relatorio.cabecalhos_com_conteudo == 1
    assert relatorio.rodapes_com_conteudo == 1


def test_auditoria_nao_modifica_documento_sprint70() -> None:
    documento = Document()

    documento.add_paragraph("Audiência em 10/01/2026")

    texto_antes = documento.paragraphs[0].text
    xml_antes = documento.element.xml

    auditor = DocumentAuditor()
    auditor.auditar(documento)

    assert documento.paragraphs[0].text == texto_antes
    assert documento.element.xml == xml_antes
def test_atualizacao_preserva_estrutura_sprint70(tmp_path) -> None:
    from pathlib import Path

    from services.text_updater import TextUpdater

    origem = Path("agenda_teste_2026.docx")

    assert origem.is_file(), (
        f"Documento de teste não encontrado: {origem}"
    )

    documento = Document(str(origem))
    auditor = DocumentAuditor()

    relatorio_antes = auditor.auditar(documento)

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    destino = tmp_path / "agenda_atualizada_sprint70.docx"

    documento.save(str(destino))

    reaberto = Document(str(destino))

    relatorio_depois = auditor.auditar(reaberto)

    textos = "\n".join(
        paragrafo.text
        for paragrafo in reaberto.paragraphs
    )

    assert quantidade == 3

    assert relatorio_antes == relatorio_depois

    assert "10/01/2027" in textos
    assert "15/05/2027" in textos
    assert "20/12/2027" in textos

    assert "10/01/2026" not in textos
    assert "15/05/2026" not in textos
    assert "20/12/2026" not in textos 
def test_atualizacao_tabelas_e_formatacao_sprint70(tmp_path) -> None:
    from services.text_updater import TextUpdater

    documento = Document()

    documento.add_paragraph("Agenda Universal PRO")

    tabela = documento.add_table(rows=2, cols=2)

    celula_1 = tabela.cell(0, 0)
    celula_1.text = ""

    run_1 = celula_1.paragraphs[0].add_run(
        "Audiência em 10/01/2026"
    )
    run_1.bold = True

    celula_2 = tabela.cell(0, 1)
    celula_2.text = "Reunião em 15/05/2026"

    celula_3 = tabela.cell(1, 0)
    celula_3.text = "Compromisso em 20/12/2026"

    celula_4 = tabela.cell(1, 1)
    celula_4.text = "Observações"

    auditor = DocumentAuditor()

    estrutura_antes = auditor.auditar(documento)

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    destino = tmp_path / "agenda_tabelas_sprint70.docx"

    documento.save(str(destino))

    reaberto = Document(str(destino))

    estrutura_depois = auditor.auditar(reaberto)

    assert quantidade == 3
    assert estrutura_antes == estrutura_depois

    tabela_atualizada = reaberto.tables[0]

    assert "10/01/2027" in tabela_atualizada.cell(0, 0).text
    assert "15/05/2027" in tabela_atualizada.cell(0, 1).text
    assert "20/12/2027" in tabela_atualizada.cell(1, 0).text

    assert tabela_atualizada.cell(1, 1).text == "Observações"

    runs_atualizados = tabela_atualizada.cell(
        0, 0
    ).paragraphs[0].runs

    runs_com_data = [
        run
        for run in runs_atualizados
        if "10/01/2027" in run.text
    ]

    assert len(runs_com_data) == 1
    assert runs_com_data[0].bold is True 