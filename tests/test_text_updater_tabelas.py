"""
Agenda Universal PRO
Arquivo: tests/test_text_updater_tabelas.py

Testa a atualização automática de datas dentro de tabelas.
"""

from pathlib import Path

from docx import Document

from core.agenda_updater import AgendaUpdater


ARQUIVO_ORIGINAL = Path("modelos") / "teste_text_updater_tabelas.docx"
ARQUIVO_ATUALIZADO = (
    Path("modelos") / "teste_text_updater_tabelas_atualizado.docx"
)


def criar_documento_teste() -> None:
    """Cria documento contendo datas dentro de uma tabela."""

    documento = Document()

    documento.add_heading(
        "Agenda Universal PRO",
        level=1,
    )

    documento.add_paragraph(
        "Documento utilizado para testar atualização em tabelas."
    )

    tabela = documento.add_table(
        rows=4,
        cols=2,
    )

    tabela.cell(0, 0).text = "Data"
    tabela.cell(0, 1).text = "Evento"

    tabela.cell(1, 0).text = "05/01/2026"
    tabela.cell(1, 1).text = "Audiência"

    tabela.cell(2, 0).text = "10/02/2026"
    tabela.cell(2, 1).text = "Reunião"

    tabela.cell(3, 0).text = "15/03/2027"
    tabela.cell(3, 1).text = "Prazo"

    documento.save(str(ARQUIVO_ORIGINAL))


def validar_tabela() -> None:
    """Verifica se as datas da tabela foram atualizadas."""

    documento = Document(
        str(ARQUIVO_ATUALIZADO)
    )

    assert len(documento.tables) == 1

    tabela = documento.tables[0]

    assert tabela.cell(1, 0).text == "05/01/2027"
    assert tabela.cell(2, 0).text == "10/02/2027"
    assert tabela.cell(3, 0).text == "15/03/2027"

    assert tabela.cell(1, 1).text == "Audiência"
    assert tabela.cell(2, 1).text == "Reunião"
    assert tabela.cell(3, 1).text == "Prazo"


def main() -> None:
    """Executa o teste de atualização em tabelas."""

    print("\n=== TESTE DO TEXTUPDATER - TABELAS ===\n")

    ARQUIVO_ORIGINAL.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    criar_documento_teste()

    updater = AgendaUpdater()

    quantidade = updater.atualizar_ano(
        ARQUIVO_ORIGINAL,
        2026,
        2027,
        ARQUIVO_ATUALIZADO,
    )

    print(
        f"Substituições realizadas: {quantidade}"
    )

    validar_tabela()

    print("Datas da tabela verificadas.")
    print("Conteúdo das células preservado.")
    print("\nTeste concluído com sucesso.")


if __name__ == "__main__":
    main()