"""
Agenda Universal PRO
Arquivo: tests/test_agenda_config_application.py

Teste de integração entre AgendaConfig e AgendaApplication.
"""

from pathlib import Path

from docx import Document

from config.agenda_config import AgendaConfig
from services.agenda_application import AgendaApplication


def criar_documento_teste(arquivo: Path) -> None:
    """Cria e salva um documento Word para o teste."""

    documento = Document()

    documento.add_paragraph("Audiência em 10/01/2026.")
    documento.add_paragraph("Reunião em 15/05/2026.")
    documento.add_paragraph("Evento em 20/12/2027.")

    tabela = documento.add_table(rows=2, cols=2)

    tabela.cell(0, 0).text = "Data"
    tabela.cell(0, 1).text = "Evento"
    tabela.cell(1, 0).text = "25/06/2026"
    tabela.cell(1, 1).text = "Reunião"

    documento.save(str(arquivo))


def validar_documento(arquivo: Path) -> None:
    """Valida o documento atualizado."""

    documento = Document(str(arquivo))

    assert documento.paragraphs[0].text == "Audiência em 10/01/2027."
    assert documento.paragraphs[1].text == "Reunião em 15/05/2027."
    assert documento.paragraphs[2].text == "Evento em 20/12/2027."

    tabela = documento.tables[0]

    assert tabela.cell(1, 0).text == "25/06/2027"


def test_agenda_config_integrada_com_agenda_application(
    tmp_path: Path,
) -> None:
    """Testa o fluxo integrado de configuração e atualização."""

    arquivo_origem = tmp_path / "agenda_2026.docx"
    arquivo_destino = tmp_path / "agenda_2027.docx"

    criar_documento_teste(arquivo_origem)

    config = AgendaConfig(
        arquivo_origem=arquivo_origem,
        arquivo_destino=arquivo_destino,
        ano_origem=2026,
        ano_destino=2027,
    )

    aplicacao = AgendaApplication()

    quantidade = aplicacao.atualizar(config)

    assert quantidade == 3
    assert arquivo_destino.exists()

    validar_documento(arquivo_destino)