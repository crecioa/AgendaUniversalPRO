"""
Agenda Universal PRO
Arquivo: tests/test_agenda_integracao_completa.py

Teste de integração completa do processamento da agenda.
"""

from pathlib import Path

from docx import Document

from core.agenda_updater import AgendaUpdater


ARQUIVO_ORIGINAL = Path("modelos") / "teste_integracao_completa.docx"
ARQUIVO_ATUALIZADO = (
    Path("modelos") / "teste_integracao_completa_atualizada.docx"
)


def criar_documento_teste() -> None:
    """Cria o documento utilizado no teste de integração."""

    documento = Document()

    documento.add_paragraph("Agenda Universal PRO")
    documento.add_paragraph("Audiência em 05/01/2026.")
    documento.add_paragraph("Reunião em 10/02/2026.")
    documento.add_paragraph("Prazo em 15/03/2027.")
    documento.add_paragraph("Texto sem nenhuma data.")

    documento.save(str(ARQUIVO_ORIGINAL))


def verificar_documento_atualizado() -> None:
    """Verifica o conteúdo do documento após a atualização."""

    documento = Document(str(ARQUIVO_ATUALIZADO))

    textos = [
        paragrafo.text
        for paragrafo in documento.paragraphs
    ]

    texto_completo = "\n".join(textos)

    assert "05/01/2027" in texto_completo
    assert "10/02/2027" in texto_completo
    assert "15/03/2027" in texto_completo

    assert "05/01/2026" not in texto_completo
    assert "10/02/2026" not in texto_completo

    assert "Texto sem nenhuma data." in texto_completo


def main() -> None:
    """Executa o teste de integração completa."""

    print("\n=== TESTE DE INTEGRAÇÃO COMPLETA ===\n")

    ARQUIVO_ORIGINAL.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    criar_documento_teste()

    updater = AgendaUpdater()

    quantidade = updater.atualizar_ano(
        ARQUIVO_ORIGINAL,
        2026,
        2027,
        ARQUIVO_ATUALIZADO,
    )

    print(f"Datas atualizadas: {quantidade}")

    assert quantidade == 2

    assert ARQUIVO_ATUALIZADO.exists()

    verificar_documento_atualizado()

    print("Documento atualizado criado com sucesso.")
    print("Conteúdo verificado com sucesso.")
    print("\nTeste de integração concluído com sucesso.")


if __name__ == "__main__":
    main()