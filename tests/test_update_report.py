"""Testes do relatório de atualização (Sprint 5.13)."""

import json
from pathlib import Path
from unittest.mock import Mock

import pytest
from docx import Document

from core.agenda_updater import AgendaUpdater
from services.update_report import UpdateReportService


def _updater() -> AgendaUpdater:
    updater = AgendaUpdater()
    updater.backup_service.criar_backup = Mock()  # type: ignore[method-assign]
    return updater


def test_relatorio_ano_com_paragrafo_e_tabela(tmp_path: Path) -> None:
    origem = tmp_path / "origem.docx"
    destino = tmp_path / "destino.docx"
    documento = Document()
    documento.add_paragraph("Audiência 05/01/2026")
    documento.add_table(rows=1, cols=1).cell(0, 0).text = "10/02/2026"
    documento.save(str(origem))
    updater = _updater()
    assert updater.atualizar_ano(origem, 2026, 2027, destino) == 2
    relatorio = updater.ultimo_relatorio
    assert relatorio is not None
    assert (relatorio.previstas, relatorio.atualizadas, relatorio.pendentes) == (2, 2, 0)
    assert relatorio.status == "sucesso"
    assert len(relatorio.alteracoes) == 2
    assert any("tabela[0]" in item.local for item in relatorio.alteracoes)


def test_relatorio_substituicoes(tmp_path: Path) -> None:
    origem = tmp_path / "origem.docx"
    destino = tmp_path / "destino.docx"
    documento = Document()
    documento.add_paragraph("05/01/2026")
    documento.save(str(origem))
    updater = _updater()
    assert updater.atualizar_agenda(origem, {"05/01/2026": "05/01/2027"}, destino) == 1
    assert updater.ultimo_relatorio is not None
    assert updater.ultimo_relatorio.alteracoes[0].datas_depois == ["05/01/2027"]


def test_exportacao_json_sem_sobrescrever(tmp_path: Path) -> None:
    origem = tmp_path / "origem.docx"
    destino = tmp_path / "destino.docx"
    relatorio_json = tmp_path / "relatorio.json"
    documento = Document()
    documento.add_paragraph("05/01/2026")
    documento.save(str(origem))
    updater = _updater()
    updater.atualizar_ano(origem, 2026, 2027, destino)
    assert updater.exportar_ultimo_relatorio(relatorio_json) == relatorio_json
    dados = json.loads(relatorio_json.read_text(encoding="utf-8"))
    assert dados["atualizadas"] == 1
    assert dados["alteracoes"][0]["datas_antes"] == ["05/01/2026"]
    with pytest.raises(FileExistsError):
        updater.exportar_ultimo_relatorio(relatorio_json)


def test_sem_relatorio_antes_de_atualizar(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Não existe relatório"):
        _updater().exportar_ultimo_relatorio(tmp_path / "relatorio.json")


def test_falha_de_validacao_nao_disponibiliza_relatorio(tmp_path: Path) -> None:
    origem = tmp_path / "origem.docx"
    destino = tmp_path / "destino.docx"
    documento = Document()
    documento.add_paragraph("29/02/2024")
    documento.save(str(origem))
    updater = _updater()
    with pytest.raises(ValueError, match="Validação de conteúdo falhou"):
        updater.atualizar_ano(origem, 2024, 2025, destino)
    assert updater.ultimo_relatorio is None
    assert not destino.exists()


def test_exportacao_exige_json(tmp_path: Path) -> None:
    relatorio = UpdateReportService.criar("a.docx", "b.docx", "ano", {}, {}, 0, 0)
    with pytest.raises(ValueError, match="extensão .json"):
        UpdateReportService.exportar(relatorio, tmp_path / "arquivo.txt")
