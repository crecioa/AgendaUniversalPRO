from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from services.text_updater import TextUpdater


def adicionar_hyperlink(paragrafo, texto: str, url: str) -> None:
    rel_id = paragrafo.part.relate_to(
        url,
        RT.HYPERLINK,
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)

    run = OxmlElement("w:r")
    elemento_texto = OxmlElement("w:t")
    elemento_texto.text = texto

    run.append(elemento_texto)
    hyperlink.append(run)

    paragrafo._p.append(hyperlink)


def test_atualizar_ano_preserva_hyperlink_sprint68() -> None:
    documento = Document()

    paragrafo = documento.add_paragraph(
        "Audiência em 10/01/2026. "
    )

    adicionar_hyperlink(
        paragrafo,
        "Consultar processo",
        "https://example.com/processo",
    )

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 1
    assert "10/01/2027" in paragrafo.text

    assert len(paragrafo.hyperlinks) == 1

    hyperlink = paragrafo.hyperlinks[0]

    assert hyperlink.address == "https://example.com/processo"
    assert hyperlink.text == "Consultar processo"


def test_data_dentro_de_hyperlink_sprint68() -> None:
    documento = Document()
    paragrafo = documento.add_paragraph("Processo: ")

    adicionar_hyperlink(
        paragrafo,
        "Consultar em 10/01/2026",
        "https://example.com/processo",
    )

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 1
    assert "10/01/2027" in paragrafo.hyperlinks[0].text
    assert (
        paragrafo.hyperlinks[0].address
        == "https://example.com/processo"
    )


def test_data_fragmentada_dentro_de_hyperlink_sprint68() -> None:
    documento = Document()
    paragrafo = documento.add_paragraph("Processo: ")

    adicionar_hyperlink(
        paragrafo,
        "Consultar em 10/01/2026",
        "https://example.com/processo",
    )

    hyperlink = paragrafo.hyperlinks[0]
    run_original = hyperlink.runs[0]
    run_original.text = "Consultar em 10/01/"

    from docx.oxml import OxmlElement

    segundo_run = OxmlElement("w:r")
    texto = OxmlElement("w:t")
    texto.text = "2026"
    segundo_run.append(texto)
    hyperlink._hyperlink.append(segundo_run)

    assert len(paragrafo.hyperlinks[0].runs) == 2

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 1
    assert paragrafo.hyperlinks[0].text == (
        "Consultar em 10/01/2027"
    )
    assert (
        paragrafo.hyperlinks[0].address
        == "https://example.com/processo"
    )


def test_hyperlink_preservado_apos_salvar_sprint68(tmp_path) -> None:
    documento = Document()
    paragrafo = documento.add_paragraph("Processo: ")

    adicionar_hyperlink(
        paragrafo,
        "Consultar em 10/01/2026",
        "https://example.com/processo",
    )

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    destino = tmp_path / "agenda_hyperlink_sprint68.docx"
    documento.save(destino)

    documento_reaberto = Document(destino)
    hyperlink = documento_reaberto.paragraphs[0].hyperlinks[0]

    assert quantidade == 1
    assert hyperlink.text == "Consultar em 10/01/2027"
    assert hyperlink.address == "https://example.com/processo"


def test_atualizacao_manual_em_hyperlink_sprint68() -> None:
    documento = Document()
    paragrafo = documento.add_paragraph("Processo: ")

    adicionar_hyperlink(
        paragrafo,
        "Audiência em 10/01/2026",
        "https://example.com/processo",
    )

    quantidade = TextUpdater()._atualizar_paragrafo(
        paragrafo,
        {"10/01/2026": "10/01/2027"},
    )

    assert quantidade == 1
    assert (
        paragrafo.hyperlinks[0].text
        == "Audiência em 10/01/2027"
    )
    assert (
        paragrafo.hyperlinks[0].address
        == "https://example.com/processo"
    )
