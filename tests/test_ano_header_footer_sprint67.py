from docx import Document

from services.text_updater import TextUpdater


def test_atualizar_ano_em_cabecalho_e_rodape_sprint67() -> None:
    documento = Document()
    secao = documento.sections[0]

    cabecalho = secao.header.paragraphs[0]
    cabecalho.add_run("Agenda ")
    data_cabecalho = cabecalho.add_run("10/01/2026")
    data_cabecalho.bold = True

    rodape = secao.footer.paragraphs[0]
    rodape.add_run("Revisão ")
    data_rodape = rodape.add_run("15/05/2026")
    data_rodape.italic = True

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 2
    assert cabecalho.text == "Agenda 10/01/2027"
    assert rodape.text == "Revisão 15/05/2027"
    assert cabecalho.runs[1].bold is True
    assert rodape.runs[1].italic is True
