
"""
Agenda Universal PRO
Arquivo: tests/test_word_service.py

Testes do serviço responsável pela manipulação de documentos Word.

Sprint 4.11:
- Testa abertura de documentos.
- Testa abertura com arquivo inexistente.
- Testa salvamento de documentos.
- Testa contagem de parágrafos.
- Testa contagem de tabelas.
- Testa registro das operações no logger.
"""

from pathlib import Path
from unittest.mock import Mock

import pytest
from typing import Any

from docx import Document

from services.word_service import WordService


def test_abrir_documento(tmp_path: Path) -> None:
    """Testa a abertura de um documento Word."""

    arquivo = tmp_path / "agenda.docx"

    documento = Document()
    documento.add_paragraph("Agenda Universal PRO")
    documento.save(str(arquivo))

    service = WordService()

    resultado = service.abrir_documento(arquivo)

    assert resultado is not None
    assert len(resultado.paragraphs) == 1
    assert resultado.paragraphs[0].text == "Agenda Universal PRO"


def test_abrir_documento_com_string(tmp_path: Path) -> None:
    """Testa a abertura utilizando caminho em formato string."""

    arquivo = tmp_path / "agenda.docx"

    documento = Document()
    documento.add_paragraph("Documento de teste")
    documento.save(str(arquivo))

    service = WordService()

    resultado = service.abrir_documento(str(arquivo))

    assert resultado is not None
    assert resultado.paragraphs[0].text == "Documento de teste"


def test_abrir_documento_inexistente(tmp_path: Path) -> None:
    """Testa erro ao tentar abrir arquivo inexistente."""

    arquivo = tmp_path / "nao_existe.docx"

    service = WordService()

    try:
        service.abrir_documento(arquivo)
        assert False, "Era esperada uma exceção FileNotFoundError."
    except FileNotFoundError as erro:
        assert "Arquivo não encontrado" in str(erro)


def test_salvar_documento(tmp_path: Path) -> None:
    """Testa o salvamento de um documento Word."""

    destino = tmp_path / "agenda_salva.docx"

    documento = Document()
    documento.add_paragraph("Documento salvo com sucesso.")

    service = WordService()

    service.salvar_documento(documento, destino)

    assert destino.exists()

    documento_salvo = Document(str(destino))

    assert len(documento_salvo.paragraphs) == 1
    assert (
        documento_salvo.paragraphs[0].text
        == "Documento salvo com sucesso."
    )


def test_salvar_documento_com_string(tmp_path: Path) -> None:
    """Testa o salvamento utilizando caminho em formato string."""

    destino = tmp_path / "agenda_string.docx"

    documento = Document()
    documento.add_paragraph("Salvamento com string.")

    service = WordService()

    service.salvar_documento(documento, str(destino))

    assert destino.exists()

    documento_salvo = Document(str(destino))

    assert (
        documento_salvo.paragraphs[0].text
        == "Salvamento com string."
    )


def test_contar_paragrafos() -> None:
    """Testa a contagem de parágrafos."""

    documento = Document()

    documento.add_paragraph("Primeiro parágrafo.")
    documento.add_paragraph("Segundo parágrafo.")
    documento.add_paragraph("Terceiro parágrafo.")

    resultado = WordService.contar_paragrafos(documento)

    assert resultado == 3


def test_contar_tabelas() -> None:
    """Testa a contagem de tabelas."""

    documento = Document()

    documento.add_table(rows=2, cols=2)
    documento.add_table(rows=3, cols=3)

    resultado = WordService.contar_tabelas(documento)

    assert resultado == 2


def test_contar_paragrafos_documento_vazio() -> None:
    """Testa a contagem de parágrafos em documento vazio."""

    documento = Document()

    resultado = WordService.contar_paragrafos(documento)

    assert resultado == 0


def test_contar_tabelas_documento_sem_tabelas() -> None:
    """Testa a contagem quando não existem tabelas."""

    documento = Document()

    resultado = WordService.contar_tabelas(documento)

    assert resultado == 0


def test_abrir_documento_registra_log(
    tmp_path: Path,
    monkeypatch: Any,
) -> None:
    """Testa o registro de log ao abrir um documento."""

    arquivo = tmp_path / "agenda.log.docx"

    documento = Document()
    documento.add_paragraph("Teste de log.")
    documento.save(str(arquivo))

    mensagens: list[str] = []

    class LoggerFalso:
        def info(
            self,
            mensagem: str,
            *args: object,
        ) -> None:
            mensagens.append(
                mensagem % args if args else mensagem
            )

    monkeypatch.setattr(
        "services.word_service.LoggerService.get_logger",
        lambda: LoggerFalso(),
    )

    service = WordService()
    service.abrir_documento(arquivo)

    assert any(
        "Abrindo documento: agenda.log.docx" in mensagem
        for mensagem in mensagens
    )


def test_salvar_documento_registra_log(
    tmp_path: Path,
    monkeypatch: Any,
) -> None:
    """Testa o registro de log ao salvar um documento."""

    destino = tmp_path / "salvo.log.docx"

    documento = Document()
    documento.add_paragraph("Teste de salvamento.")

    mensagens: list[str] = []

    class LoggerFalso:
        def info(
            self,
            mensagem: str,
            *args: object,
        ) -> None:
            mensagens.append(
                mensagem % args if args else mensagem
            )

    monkeypatch.setattr(
        "services.word_service.LoggerService.get_logger",
        lambda: LoggerFalso(),
    )

    service = WordService()
    service.salvar_documento(documento, destino)

    assert destino.exists()

    assert any(
        "Documento salvo: salvo.log.docx" in mensagem
        for mensagem in mensagens
    )


def test_salvar_documento_recusa_destino_existente(tmp_path: Path) -> None:
    """Não substitui documento já existente."""
    destino = tmp_path / "existente.docx"
    destino.write_bytes(b"conteudo original")

    service = WordService()
    documento = Document()

    with pytest.raises(FileExistsError):
        service.salvar_documento(documento, destino)

    assert destino.read_bytes() == b"conteudo original"
    assert list(tmp_path.iterdir()) == [destino]


def test_salvar_documento_limpa_temporario_em_falha(tmp_path: Path) -> None:
    """Falha na escrita não publica documento parcial."""
    destino = tmp_path / "resultado.docx"
    documento = Mock()
    documento.save.side_effect = OSError("falha simulada")

    with pytest.raises(OSError, match="falha simulada"):
        WordService().salvar_documento(documento, destino)

    assert not destino.exists()
    assert list(tmp_path.iterdir()) == []


def test_salvar_documento_protege_destino_criado_durante_gravacao(
    tmp_path: Path,
) -> None:
    """Não substitui arquivo criado por outro processo durante a escrita."""
    destino = tmp_path / "concorrente.docx"

    class DocumentoConcorrente:
        def save(self, caminho: str) -> None:
            Document().save(caminho)
            destino.write_bytes(b"arquivo de outro processo")

    with pytest.raises(FileExistsError):
        WordService().salvar_documento(
            DocumentoConcorrente(),  # type: ignore[arg-type]
            destino,
        )

    assert destino.read_bytes() == b"arquivo de outro processo"
    assert list(tmp_path.iterdir()) == [destino]


def test_salvar_documento_recusa_pacote_corrompido(tmp_path: Path) -> None:
    """Não publica um arquivo ZIP inválido após o save."""
    destino = tmp_path / "corrompido.docx"
    documento = Mock()
    documento.save.side_effect = lambda arquivo: Path(arquivo).write_bytes(
        b"nao e um documento Word"
    )

    with pytest.raises(ValueError, match="inválido ou está corrompido"):
        WordService().salvar_documento(documento, destino)

    assert not destino.exists()
    assert list(tmp_path.iterdir()) == []


def test_salvar_documento_recusa_zip_sem_documento_word(
    tmp_path: Path,
) -> None:
    """Um ZIP válido sem estrutura DOCX também é recusado."""
    from zipfile import ZipFile

    destino = tmp_path / "zip_invalido.docx"

    class DocumentoZipInvalido:
        def save(self, arquivo: str) -> None:
            with ZipFile(arquivo, "w") as pacote:
                pacote.writestr("arquivo.txt", "sem estrutura Word")

    with pytest.raises(ValueError, match="inválido ou está corrompido"):
        WordService().salvar_documento(
            DocumentoZipInvalido(),  # type: ignore[arg-type]
            destino,
        )

    assert not destino.exists()
    assert list(tmp_path.iterdir()) == []
