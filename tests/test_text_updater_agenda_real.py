from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from core.date_parser import DateParser
from services.text_updater import TextUpdater


def criar_documento_real() -> Path:
    """Cria um documento com os padrões reais da agenda."""
    with TemporaryDirectory() as diretorio:
        caminho = Path(diretorio) / "agenda_2026.docx"

        documento = Document()

        documento.add_paragraph(
            "AUDIÊNCIAS – Quinta-feira, 01/01/2026"
        )
        documento.add_paragraph(
            "ACORDO – Quinta-feira, 01/01/2026"
        )
        documento.add_paragraph(
            "AUDIÊNCIAS – Segunda-feira, 05/01/2026"
        )
        documento.add_paragraph(
            "ACORDO – Segunda-feira, 05/01/2026"
        )

        documento.save(str(caminho))

        return caminho


def test_atualizar_padrao_audiencias() -> None:
    """Testa atualização do padrão AUDIÊNCIAS."""
    with TemporaryDirectory() as diretorio:
        caminho = Path(diretorio) / "agenda_2026.docx"
        destino = Path(diretorio) / "agenda_2027.docx"

        documento = Document()
        documento.add_paragraph(
            "AUDIÊNCIAS – Quinta-feira, 01/01/2026"
        )

        documento.save(str(caminho))

        documento = Document(str(caminho))

        updater = TextUpdater()

        quantidade = updater.atualizar_ano(
            documento,
            2026,
            2027,
        )

        documento.save(str(destino))

        assert quantidade == 1
        assert documento.paragraphs[0].text == (
            "AUDIÊNCIAS – Sexta-feira, 01/01/2027"
        )


def test_atualizar_padrao_acordo() -> None:
    """Testa atualização do padrão ACORDO."""
    with TemporaryDirectory() as diretorio:
        caminho = Path(diretorio) / "agenda_2026.docx"
        destino = Path(diretorio) / "agenda_2027.docx"

        documento = Document()
        documento.add_paragraph(
            "ACORDO – Quinta-feira, 01/01/2026"
        )

        documento.save(str(caminho))

        documento = Document(str(caminho))

        updater = TextUpdater()

        quantidade = updater.atualizar_ano(
            documento,
            2026,
            2027,
        )

        documento.save(str(destino))

        assert quantidade == 1
        assert documento.paragraphs[0].text == (
            "ACORDO – Sexta-feira, 01/01/2027"
        )


def test_atualizar_datas_reais() -> None:
    """Testa várias datas reais da agenda."""
    documento = Document()

    documento.add_paragraph(
        "AUDIÊNCIAS – Quinta-feira, 01/01/2026"
    )
    documento.add_paragraph(
        "ACORDO – Quinta-feira, 01/01/2026"
    )
    documento.add_paragraph(
        "AUDIÊNCIAS – Segunda-feira, 05/01/2026"
    )
    documento.add_paragraph(
        "ACORDO – Segunda-feira, 05/01/2026"
    )
    documento.add_paragraph(
        "AUDIÊNCIAS – Terça-feira, 10/02/2026"
    )
    documento.add_paragraph(
        "AUDIÊNCIAS – Domingo, 15/03/2026"
    )

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 6

    assert documento.paragraphs[0].text == (
        "AUDIÊNCIAS – Sexta-feira, 01/01/2027"
    )
    assert documento.paragraphs[1].text == (
        "ACORDO – Sexta-feira, 01/01/2027"
    )
    assert documento.paragraphs[2].text == (
        "AUDIÊNCIAS – Terça-feira, 05/01/2027"
    )
    assert documento.paragraphs[3].text == (
        "ACORDO – Terça-feira, 05/01/2027"
    )
    assert documento.paragraphs[4].text == (
        "AUDIÊNCIAS – Quarta-feira, 10/02/2027"
    )
    assert documento.paragraphs[5].text == (
        "AUDIÊNCIAS – Segunda-feira, 15/03/2027"
    )


def test_atualizar_datas_reais_em_arquivo() -> None:
    """Testa atualização usando arquivo DOCX."""
    with TemporaryDirectory() as diretorio:
        origem = Path(diretorio) / "agenda_2026.docx"
        destino = Path(diretorio) / "agenda_2027.docx"

        documento = Document()

        documento.add_paragraph(
            "AUDIÊNCIAS – Quinta-feira, 01/01/2026"
        )
        documento.add_paragraph(
            "ACORDO – Quinta-feira, 01/01/2026"
        )
        documento.add_paragraph(
            "AUDIÊNCIAS – Segunda-feira, 05/01/2026"
        )
        documento.add_paragraph(
            "ACORDO – Segunda-feira, 05/01/2026"
        )

        documento.save(str(origem))

        documento = Document(str(origem))

        updater = TextUpdater()

        quantidade = updater.atualizar_ano(
            documento,
            2026,
            2027,
        )

        documento.save(str(destino))

        assert quantidade == 4
        assert destino.exists()

        documento_final = Document(str(destino))

        assert documento_final.paragraphs[0].text == (
            "AUDIÊNCIAS – Sexta-feira, 01/01/2027"
        )
        assert documento_final.paragraphs[1].text == (
            "ACORDO – Sexta-feira, 01/01/2027"
        )
        assert documento_final.paragraphs[2].text == (
            "AUDIÊNCIAS – Terça-feira, 05/01/2027"
        )
        assert documento_final.paragraphs[3].text == (
            "ACORDO – Terça-feira, 05/01/2027"
        )


def test_date_parser_dias_reais() -> None:
    """Confirma os dias da semana das datas usadas no teste."""
    assert (
        DateParser.obter_dia_semana("01/01/2026")
        == "Quinta-feira"
    )
    assert (
        DateParser.obter_dia_semana("01/01/2027")
        == "Sexta-feira"
    )
    assert (
        DateParser.obter_dia_semana("05/01/2026")
        == "Segunda-feira"
    )
    assert (
        DateParser.obter_dia_semana("05/01/2027")
        == "Terça-feira"
    )
    assert (
        DateParser.obter_dia_semana("10/02/2026")
        == "Terça-feira"
    )
    assert (
        DateParser.obter_dia_semana("10/02/2027")
        == "Quarta-feira"
    )
    assert (
        DateParser.obter_dia_semana("15/03/2026")
        == "Domingo"
    )
    assert (
        DateParser.obter_dia_semana("15/03/2027")
        == "Segunda-feira"
    )