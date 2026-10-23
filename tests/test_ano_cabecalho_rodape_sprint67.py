from docx import Document

from services.text_updater import TextUpdater


def test_atualizar_ano_preserva_formatacao_sprint67() -> None:
    documento = Document()
    paragrafo = documento.add_paragraph()

    titulo = paragrafo.add_run("Audiência: ")
    titulo.bold = True

    data = paragrafo.add_run("10/01/2026")
    data.italic = True

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 1
    assert paragrafo.text == "Audiência: 10/01/2027"
    assert paragrafo.runs[0].bold is True
    assert paragrafo.runs[1].italic is True


def test_atualizar_datas_em_cabecalho_e_rodape_sprint67() -> None:
    documento = Document()
    secao = documento.sections[0]

    cabecalho = secao.header.paragraphs[0]
    cabecalho.add_run("Agenda 10/01/2026")

    rodape = secao.footer.paragraphs[0]
    rodape.add_run("Revisão 15/05/2026")

    quantidade = TextUpdater().atualizar_datas(
        documento,
        {
            "10/01/2026": "10/01/2027",
            "15/05/2026": "15/05/2027",
        },
    )

    assert quantidade == 2
    assert cabecalho.text == "Agenda 10/01/2027"
    assert rodape.text == "Revisão 15/05/2027"
