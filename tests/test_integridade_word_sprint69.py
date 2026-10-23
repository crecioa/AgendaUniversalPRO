from docx import Document

from services.text_updater import TextUpdater
from tests.test_hyperlinks_sprint68 import adicionar_hyperlink


def test_multiplos_hyperlinks_mesmo_paragrafo_sprint69() -> None:
    documento = Document()
    paragrafo = documento.add_paragraph("Processos: ")

    adicionar_hyperlink(
        paragrafo,
        "Audiência em 10/01/2026",
        "https://example.com/processo-1",
    )

    paragrafo.add_run(" | ")

    adicionar_hyperlink(
        paragrafo,
        "Prazo em 20/02/2026",
        "https://example.com/processo-2",
    )

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    hyperlinks = paragrafo.hyperlinks

    assert quantidade == 2
    assert len(hyperlinks) == 2

    assert "10/01/2027" in hyperlinks[0].text
    assert "20/02/2027" in hyperlinks[1].text

    assert (
        hyperlinks[0].address
        == "https://example.com/processo-1"
    )
    assert (
        hyperlinks[1].address
        == "https://example.com/processo-2"
    )

    assert " | " in paragrafo.text


def test_data_dividida_entre_texto_e_hyperlink_sprint69() -> None:
    documento = Document()
    paragrafo = documento.add_paragraph("Audiência em 10/01/")

    adicionar_hyperlink(
        paragrafo,
        "2026",
        "https://example.com/processo",
    )

    texto_original = paragrafo.text
    endereco_original = paragrafo.hyperlinks[0].address

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 1
    assert paragrafo.text == "Audiência em 10/01/2027"
    assert paragrafo.hyperlinks[0].text == "2027"
    assert paragrafo.hyperlinks[0].address == endereco_original
    assert texto_original == "Audiência em 10/01/2026"


def test_formatacao_mista_preservada_sprint69(tmp_path) -> None:
    documento = Document()
    paragrafo = documento.add_paragraph()

    run_inicial = paragrafo.add_run("Audiência em 10/01/")
    run_inicial.bold = True

    adicionar_hyperlink(
        paragrafo,
        "2026",
        "https://example.com/processo",
    )

    hyperlink = paragrafo.hyperlinks[0]
    hyperlink.runs[0].italic = True

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 1
    assert paragrafo.text == "Audiência em 10/01/2027"
    assert run_inicial.bold is True
    assert hyperlink.runs[0].italic is True
    assert hyperlink.text == "2027"

    destino = tmp_path / "formatacao_mista_sprint69.docx"
    documento.save(destino)

    reaberto = Document(destino)
    paragrafo_reaberto = reaberto.paragraphs[0]
    hyperlink_reaberto = paragrafo_reaberto.hyperlinks[0]

    assert paragrafo_reaberto.text == "Audiência em 10/01/2027"
    assert paragrafo_reaberto.runs[0].bold is True
    assert hyperlink_reaberto.runs[0].italic is True
    assert (
        hyperlink_reaberto.address
        == "https://example.com/processo"
    )


def test_campo_page_preservado_sprint69(tmp_path) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    documento = Document()
    paragrafo = documento.add_paragraph("Agenda de 10/01/2026 - Página ")

    run = paragrafo.add_run()

    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")

    instrucao = OxmlElement("w:instrText")
    instrucao.text = " PAGE "

    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")

    resultado = OxmlElement("w:t")
    resultado.text = "1"

    fim = OxmlElement("w:fldChar")
    fim.set(qn("w:fldCharType"), "end")

    for elemento in (
        inicio,
        instrucao,
        separador,
        resultado,
        fim,
    ):
        run._r.append(elemento)

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    destino = tmp_path / "campo_page_sprint69.docx"
    documento.save(destino)

    reaberto = Document(destino)
    xml = reaberto.paragraphs[0]._p.xml

    assert quantidade == 1
    assert "Agenda de 10/01/2027" in reaberto.paragraphs[0].text
    assert " PAGE " in xml
    assert 'w:fldCharType="begin"' in xml
    assert 'w:fldCharType="separate"' in xml
    assert 'w:fldCharType="end"' in xml
def test_campo_page_em_multiplos_runs_sprint69(tmp_path) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    documento = Document()
    paragrafo = documento.add_paragraph(
        "Audiência em 10/01/2026 - Página "
    )

    def adicionar_campo(tipo: str) -> None:
        run = paragrafo.add_run()
        elemento = OxmlElement("w:fldChar")
        elemento.set(qn("w:fldCharType"), tipo)
        run._r.append(elemento)

    adicionar_campo("begin")

    run_instrucao = paragrafo.add_run()
    instrucao = OxmlElement("w:instrText")
    instrucao.text = " PAGE "
    run_instrucao._r.append(instrucao)

    adicionar_campo("separate")

    paragrafo.add_run("1")

    adicionar_campo("end")

    quantidade = TextUpdater().atualizar_ano(
        documento,
        2026,
        2027,
    )

    destino = tmp_path / "campo_page_multiplos_runs.docx"
    documento.save(destino)

    reaberto = Document(destino)
    paragrafo_reaberto = reaberto.paragraphs[0]
    xml = paragrafo_reaberto._p.xml

    assert quantidade == 1
    assert "10/01/2027" in paragrafo_reaberto.text

    assert " PAGE " in xml
    assert xml.count('w:fldCharType="begin"') == 1
    assert xml.count('w:fldCharType="separate"') == 1
    assert xml.count('w:fldCharType="end"') == 1

