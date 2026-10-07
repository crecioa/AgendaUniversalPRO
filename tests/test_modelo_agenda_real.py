from pathlib import Path
from zipfile import ZipFile

from docx import Document

from core.date_parser import DateParser


ARQUIVO = Path("modelos") / "1º Semestre 2027.docx"


def texto_documento() -> str:
    """Extrai o texto XML completo do documento Word."""
    with ZipFile(ARQUIVO) as arquivo_zip:
        xml = arquivo_zip.read("word/document.xml")

    return xml.decode("utf-8")


def texto_legivel_documento() -> str:
    """Obtém o texto visível usando python-docx."""
    documento = Document(str(ARQUIVO))

    textos = [paragrafo.text for paragrafo in documento.paragraphs]

    for tabela in documento.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                textos.append(celula.text)

    return "\n".join(textos)


def test_modelo_existe() -> None:
    """Verifica se o modelo real existe."""
    assert ARQUIVO.exists()
    assert ARQUIVO.is_file()


def test_modelo_abre() -> None:
    """Verifica se o Word consegue ser aberto."""
    documento = Document(str(ARQUIVO))

    assert len(documento.paragraphs) > 0


def test_modelo_contem_meses() -> None:
    """Verifica os seis meses do primeiro semestre."""
    texto = texto_documento()

    for mes in (
        "JANEIRO",
        "FEVEREIRO",
        "MARÇO",
        "ABRIL",
        "MAIO",
        "JUNHO",
    ):
        assert mes in texto


def test_modelo_contem_secoes_principais() -> None:
    """Verifica as seções principais da agenda."""
    texto = texto_documento()

    assert "AUDIÊNCIAS" in texto
    assert "ACORDO" in texto


def test_modelo_contem_campos_da_agenda() -> None:
    """Verifica os campos de preenchimento."""
    texto = texto_legivel_documento()

    for campo in (
        "Autor:",
        "Réu:",
        "Proc. N°:",
        "Vara:",
        "Tipo:",
        "Hora:",
        "Pasta:",
    ):
        assert campo in texto


def test_datas_do_modelo_sao_2027() -> None:
    """Verifica se as datas encontradas pertencem a 2027."""
    texto = texto_documento()

    datas = DateParser.encontrar_datas(texto)

    assert datas
    assert all(data.endswith("/2027") for data in datas)


def test_dias_da_semana_estao_corretos() -> None:
    """Verifica se as datas possuem dias da semana válidos."""
    texto = texto_documento()

    datas = DateParser.encontrar_datas(texto)

    assert datas

    for data in datas:
        dia = DateParser.obter_dia_semana(data)

        assert dia is not None


def test_primeira_data_do_modelo() -> None:
    """Verifica a primeira data da agenda."""
    texto = texto_documento()

    assert "Sexta-feira, 01/01/2027" in texto
