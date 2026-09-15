"""
Agenda Universal PRO

Arquivo: tests/test_text_updater.py

Testes do serviço TextUpdater.

Sprint 4.14:

* Testa localização de datas em parágrafos.
* Testa substituição de datas em textos.
* Testa atualização de parágrafos.
* Testa atualização de células de tabelas.
* Testa atualização de datas em documentos.
* Testa atualização automática de ano.
* Testa caminhos defensivos da atualização automática.
"""

from docx import Document

from core.date_parser import DateParser
from services.text_updater import TextUpdater


def test_localizar_datas_em_paragrafos() -> None:
    """Testa localização de datas em parágrafos."""
    documento = Document()
    documento.add_paragraph("Audiência em 05/01/2027.")
    documento.add_paragraph("Reunião em 10/02/2027.")
    documento.add_paragraph("Texto sem data.")

    updater = TextUpdater()

    ocorrencias = updater.localizar_datas(documento)

    assert len(ocorrencias) == 2
    assert ocorrencias[0] == (
        0,
        "Audiência em 05/01/2027.",
        ["05/01/2027"],
    )
    assert ocorrencias[1] == (
        1,
        "Reunião em 10/02/2027.",
        ["10/02/2027"],
    )


def test_localizar_datas_ignora_paragrafo_vazio() -> None:
    """Testa que parágrafos vazios são ignorados."""
    documento = Document()
    documento.add_paragraph("")
    documento.add_paragraph("Evento em 05/01/2027.")

    updater = TextUpdater()

    ocorrencias = updater.localizar_datas(documento)

    assert len(ocorrencias) == 1
    assert ocorrencias[0][0] == 1


def test_localizar_datas_sem_ocorrencias() -> None:
    """Testa documento sem datas."""
    documento = Document()
    documento.add_paragraph("Documento sem datas.")

    updater = TextUpdater()

    ocorrencias = updater.localizar_datas(documento)

    assert ocorrencias == []


def test_substituir_texto() -> None:
    """Testa substituição simples de texto."""
    texto, quantidade = TextUpdater._substituir_texto(
        "Audiência em 05/01/2027.",
        {"05/01/2027": "05/01/2028"},
    )

    assert texto == "Audiência em 05/01/2028."
    assert quantidade == 1


def test_substituir_texto_multiplas_ocorrencias() -> None:
    """Testa múltiplas ocorrências da mesma data."""
    texto, quantidade = TextUpdater._substituir_texto(
        "05/01/2027 e 05/01/2027.",
        {"05/01/2027": "05/01/2028"},
    )

    assert texto == "05/01/2028 e 05/01/2028."
    assert quantidade == 2


def test_substituir_texto_sem_ocorrencias() -> None:
    """Testa texto sem ocorrência."""
    texto, quantidade = TextUpdater._substituir_texto(
        "Texto sem data.",
        {"05/01/2027": "05/01/2028"},
    )

    assert texto == "Texto sem data."
    assert quantidade == 0


def test_atualizar_paragrafo() -> None:
    """Testa atualização de um parágrafo."""
    documento = Document()
    paragrafo = documento.add_paragraph(
        "Audiência em 05/01/2027."
    )

    updater = TextUpdater()

    quantidade = updater._atualizar_paragrafo(
        paragrafo,
        {"05/01/2027": "05/01/2028"},
    )

    assert quantidade == 1
    assert paragrafo.text == "Audiência em 05/01/2028."


def test_atualizar_paragrafo_sem_alteracao() -> None:
    """Testa parágrafo sem data correspondente."""
    documento = Document()
    paragrafo = documento.add_paragraph(
        "Audiência em 05/01/2027."
    )

    updater = TextUpdater()

    quantidade = updater._atualizar_paragrafo(
        paragrafo,
        {"10/02/2027": "10/02/2028"},
    )

    assert quantidade == 0
    assert paragrafo.text == "Audiência em 05/01/2027."


def test_atualizar_paragrafo_vazio() -> None:
    """Testa parágrafo vazio."""
    documento = Document()
    paragrafo = documento.add_paragraph("")

    updater = TextUpdater()

    quantidade = updater._atualizar_paragrafo(
        paragrafo,
        {"05/01/2027": "05/01/2028"},
    )

    assert quantidade == 0
    assert paragrafo.text == ""


def test_atualizar_datas_em_paragrafos() -> None:
    """Testa atualização de datas em parágrafos."""
    documento = Document()
    documento.add_paragraph("Audiência em 05/01/2027.")
    documento.add_paragraph("Reunião em 10/02/2027.")
    documento.add_paragraph("Prazo em 15/03/2027.")

    updater = TextUpdater()

    quantidade = updater.atualizar_datas(
        documento,
        {
            "05/01/2027": "05/01/2028",
            "10/02/2027": "10/02/2028",
            "15/03/2027": "15/03/2028",
        },
    )

    assert quantidade == 3
    assert documento.paragraphs[0].text == (
        "Audiência em 05/01/2028."
    )
    assert documento.paragraphs[1].text == (
        "Reunião em 10/02/2028."
    )
    assert documento.paragraphs[2].text == (
        "Prazo em 15/03/2028."
    )


def test_atualizar_datas_em_tabela() -> None:
    """Testa atualização de datas em tabela."""
    documento = Document()

    tabela = documento.add_table(
        rows=2,
        cols=2,
    )

    tabela.cell(0, 0).text = "Data"
    tabela.cell(0, 1).text = "Evento"
    tabela.cell(1, 0).text = "05/01/2027"
    tabela.cell(1, 1).text = "Audiência"

    updater = TextUpdater()

    quantidade = updater.atualizar_datas(
        documento,
        {"05/01/2027": "05/01/2028"},
    )

    assert quantidade == 1
    assert tabela.cell(1, 0).text == "05/01/2028"
    assert tabela.cell(1, 1).text == "Audiência"


def test_atualizar_datas_paragrafos_e_tabelas() -> None:
    """Testa atualização simultânea de parágrafos e tabelas."""
    documento = Document()
    documento.add_paragraph("Audiência em 05/01/2027.")

    tabela = documento.add_table(
        rows=1,
        cols=1,
    )
    tabela.cell(0, 0).text = "Reunião em 10/02/2027."

    updater = TextUpdater()

    quantidade = updater.atualizar_datas(
        documento,
        {
            "05/01/2027": "05/01/2028",
            "10/02/2027": "10/02/2028",
        },
    )

    assert quantidade == 2
    assert documento.paragraphs[0].text == (
        "Audiência em 05/01/2028."
    )
    assert tabela.cell(0, 0).text == (
        "Reunião em 10/02/2028."
    )


def test_atualizar_datas_sem_ocorrencias() -> None:
    """Testa atualização quando nenhuma data é encontrada."""
    documento = Document()
    documento.add_paragraph("Texto sem data.")

    updater = TextUpdater()

    quantidade = updater.atualizar_datas(
        documento,
        {"05/01/2027": "05/01/2028"},
    )

    assert quantidade == 0
    assert documento.paragraphs[0].text == "Texto sem data."


def test_atualizar_ano_em_paragrafos() -> None:
    """Testa atualização automática do ano em parágrafos."""
    documento = Document()
    documento.add_paragraph("Audiência em 05/01/2026.")
    documento.add_paragraph("Reunião em 10/02/2026.")
    documento.add_paragraph("Prazo em 15/03/2027.")

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 2
    assert documento.paragraphs[0].text == (
        "Audiência em 05/01/2027."
    )
    assert documento.paragraphs[1].text == (
        "Reunião em 10/02/2027."
    )
    assert documento.paragraphs[2].text == (
        "Prazo em 15/03/2027."
    )


def test_atualizar_ano_ignora_ano_diferente() -> None:
    """Testa que datas de outro ano permanecem inalteradas."""
    documento = Document()
    documento.add_paragraph("Evento em 05/01/2027.")

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2028,
    )

    assert quantidade == 0
    assert documento.paragraphs[0].text == (
        "Evento em 05/01/2027."
    )


def test_atualizar_ano_multiplas_ocorrencias() -> None:
    """Testa múltiplas ocorrências da mesma data."""
    documento = Document()
    documento.add_paragraph(
        "05/01/2026 e novamente 05/01/2026."
    )

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 2
    assert documento.paragraphs[0].text == (
        "05/01/2027 e novamente 05/01/2027."
    )


def test_atualizar_ano_em_tabela() -> None:
    """Testa atualização automática de ano em tabela."""
    documento = Document()

    tabela = documento.add_table(
        rows=2,
        cols=2,
    )

    tabela.cell(0, 0).text = "Data"
    tabela.cell(0, 1).text = "Evento"
    tabela.cell(1, 0).text = "05/01/2026"
    tabela.cell(1, 1).text = "Audiência"

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 1
    assert tabela.cell(1, 0).text == "05/01/2027"
    assert tabela.cell(1, 1).text == "Audiência"


def test_atualizar_ano_paragrafo_vazio() -> None:
    """Testa que parágrafos vazios são ignorados."""
    documento = Document()
    documento.add_paragraph("")
    documento.add_paragraph(
        "Evento em 05/01/2026."
    )

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 1
    assert documento.paragraphs[0].text == ""
    assert documento.paragraphs[1].text == (
        "Evento em 05/01/2027."
    )


def test_atualizar_ano_sem_datas() -> None:
    """Testa documento sem datas."""
    documento = Document()
    documento.add_paragraph(
        "Documento sem datas."
    )

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 0
    assert documento.paragraphs[0].text == (
        "Documento sem datas."
    )


def test_atualizar_ano_ignora_data_invalida(
    monkeypatch,
) -> None:
    """Testa quando uma data não pode ser convertida."""
    documento = Document()
    documento.add_paragraph(
        "Evento em 05/01/2026."
    )

    monkeypatch.setattr(
        DateParser,
        "converter_data",
        staticmethod(lambda data: None),
    )

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 0
    assert documento.paragraphs[0].text == (
        "Evento em 05/01/2026."
    )


def test_atualizar_ano_ignora_data_sem_nova_data(
    monkeypatch,
) -> None:
    """Testa quando não existe nova data."""
    documento = Document()
    documento.add_paragraph(
        "Evento em 05/01/2026."
    )

    monkeypatch.setattr(
        DateParser,
        "atualizar_ano",
        staticmethod(lambda data, ano: None),
    )

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 0
    assert documento.paragraphs[0].text == (
        "Evento em 05/01/2026."
    )


def test_atualizar_ano_ignora_paragrafo_vazio_em_tabela() -> None:
    """Testa parágrafo vazio dentro de célula de tabela."""
    documento = Document()

    tabela = documento.add_table(
        rows=1,
        cols=1,
    )

    tabela.cell(0, 0).text = ""

    updater = TextUpdater()

    quantidade = updater.atualizar_ano(
        documento,
        2026,
        2027,
    )

    assert quantidade == 0
    assert tabela.cell(0, 0).text == ""