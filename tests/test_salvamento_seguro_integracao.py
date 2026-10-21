
"""Testes de integração do salvamento seguro de agendas (Sprint 5.10)."""

from pathlib import Path
from unittest.mock import Mock

import pytest
from docx import Document
from docx.document import Document as DocxDocument

from core.agenda_updater import AgendaUpdater


def criar_agenda(caminho: Path) -> None:
    """Cria um documento Word real com parágrafos e tabela."""
    documento = Document()
    documento.add_heading("Agenda Universal PRO", level=1)
    documento.add_paragraph("Audiência em 05/01/2026.")
    documento.add_paragraph("Reunião em 10/02/2026.")
    tabela = documento.add_table(rows=2, cols=2)
    tabela.cell(0, 0).text = "Data"
    tabela.cell(0, 1).text = "Evento"
    tabela.cell(1, 0).text = "15/03/2027"
    tabela.cell(1, 1).text = "Prazo"
    documento.save(str(caminho))


def criar_updater_sem_backup() -> AgendaUpdater:
    """Mantém serviços reais de Word e texto, isolando apenas o backup."""
    updater = AgendaUpdater()
    updater.backup_service.criar_backup = Mock()  # type: ignore[method-assign]
    return updater


def temporarios(caminho: Path) -> list[Path]:
    """Encontra arquivos temporários usados pelo WordService."""
    return list(caminho.parent.glob(f".{caminho.stem}.*.docx"))


def test_integracao_salva_agenda_e_preserva_original(tmp_path: Path) -> None:
    origem = tmp_path / "agenda.docx"
    destino = tmp_path / "agenda_2027.docx"
    criar_agenda(origem)
    conteudo_original = origem.read_bytes()

    quantidade = criar_updater_sem_backup().atualizar_ano(
        origem, 2026, 2027, destino
    )

    assert quantidade == 2
    assert origem.read_bytes() == conteudo_original
    assert destino.is_file()
    documento = Document(str(destino))
    texto = "\n".join(paragrafo.text for paragrafo in documento.paragraphs)
    assert "05/01/2027" in texto
    assert "10/02/2027" in texto
    assert "05/01/2026" not in texto
    assert "10/02/2026" not in texto
    assert len(documento.tables) == 1
    assert documento.tables[0].cell(1, 0).text == "15/03/2027"
    assert not temporarios(destino)


def test_integracao_recusa_destino_existente(tmp_path: Path) -> None:
    origem = tmp_path / "agenda.docx"
    destino = tmp_path / "agenda_2027.docx"
    criar_agenda(origem)
    criar_agenda(destino)
    original = origem.read_bytes()
    destino_anterior = destino.read_bytes()

    with pytest.raises(FileExistsError, match="destino já existe"):
        criar_updater_sem_backup().atualizar_ano(
            origem, 2026, 2027, destino
        )

    assert origem.read_bytes() == original
    assert destino.read_bytes() == destino_anterior
    assert not temporarios(destino)


def test_integracao_falha_no_save_limpa_temporario(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    origem = tmp_path / "agenda.docx"
    destino = tmp_path / "agenda_2027.docx"
    criar_agenda(origem)
    original = origem.read_bytes()

    def falhar_ao_salvar(self: DocxDocument, arquivo: str) -> None:
        Path(arquivo).write_bytes(b"gravacao incompleta")
        raise OSError("falha simulada na gravação")

    monkeypatch.setattr(DocxDocument, "save", falhar_ao_salvar)

    with pytest.raises(OSError, match="falha simulada"):
        criar_updater_sem_backup().atualizar_ano(
            origem, 2026, 2027, destino
        )

    assert origem.read_bytes() == original
    assert not destino.exists()
    assert not temporarios(destino)


def test_integracao_destino_criado_durante_publicacao(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    origem = tmp_path / "agenda.docx"
    destino = tmp_path / "agenda_2027.docx"
    criar_agenda(origem)
    original = origem.read_bytes()
    concorrente = b"arquivo criado por outro processo"

    def simular_concorrencia(temporario: object, caminho: object) -> None:
        assert Path(temporario).is_file()  # type: ignore[arg-type]
        Path(caminho).write_bytes(concorrente)  # type: ignore[arg-type]
        raise FileExistsError("destino criado durante a publicação")

    monkeypatch.setattr("services.word_service.os.link", simular_concorrencia)

    with pytest.raises(FileExistsError, match="destino criado"):
        criar_updater_sem_backup().atualizar_ano(
            origem, 2026, 2027, destino
        )

    assert origem.read_bytes() == original
    assert destino.read_bytes() == concorrente
    assert not temporarios(destino)


def test_integracao_nao_publica_documento_corrompido(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A atualização completa não publica um DOCX inválido."""
    origem = tmp_path / "agenda.docx"
    destino = tmp_path / "agenda_2027.docx"
    criar_agenda(origem)
    original = origem.read_bytes()

    def gravar_corrompido(self: DocxDocument, arquivo: str) -> None:
        Path(arquivo).write_bytes(b"arquivo Word corrompido")

    monkeypatch.setattr(DocxDocument, "save", gravar_corrompido)

    with pytest.raises(ValueError, match="inválido ou está corrompido"):
        criar_updater_sem_backup().atualizar_ano(
            origem, 2026, 2027, destino
        )

    assert origem.read_bytes() == original
    assert not destino.exists()
    assert not temporarios(destino)
