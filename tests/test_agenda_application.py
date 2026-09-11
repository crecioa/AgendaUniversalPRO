"""
Agenda Universal PRO
Arquivo: tests/test_agenda_application.py

Testes da camada AgendaApplication.

Sprint 4.7:
- Valida o fluxo completo de atualização de ano.
- Valida propagação de erros.
- Valida registro de exceções no logger.
"""

from pathlib import Path
from typing import Any, cast
from unittest.mock import Mock

import pytest
from docx import Document

from config.agenda_config import AgendaConfig
from services.agenda_application import AgendaApplication


def criar_documento_teste(arquivo: Path) -> None:
    """Cria um documento Word para os testes."""

    documento = Document()

    documento.add_paragraph(
        "Audiência em 05/01/2026."
    )

    documento.add_paragraph(
        "Reunião em 10/02/2026."
    )

    documento.add_paragraph(
        "Prazo em 15/03/2027."
    )

    documento.add_paragraph(
        "Texto sem data."
    )

    tabela = documento.add_table(
        rows=2,
        cols=2,
    )

    tabela.cell(0, 0).text = "Data"
    tabela.cell(0, 1).text = "Evento"
    tabela.cell(1, 0).text = "20/04/2026"
    tabela.cell(1, 1).text = "Sessão"

    documento.save(str(arquivo))


def validar_documento(arquivo: Path) -> None:
    """Valida o conteúdo do documento atualizado."""

    documento = Document(str(arquivo))

    assert (
        documento.paragraphs[0].text
        == "Audiência em 05/01/2027."
    )

    assert (
        documento.paragraphs[1].text
        == "Reunião em 10/02/2027."
    )

    assert (
        documento.paragraphs[2].text
        == "Prazo em 15/03/2027."
    )

    assert (
        documento.paragraphs[3].text
        == "Texto sem data."
    )

    assert len(documento.tables) == 1

    tabela = documento.tables[0]

    assert tabela.cell(0, 0).text == "Data"
    assert tabela.cell(0, 1).text == "Evento"
    assert tabela.cell(1, 0).text == "20/04/2027"
    assert tabela.cell(1, 1).text == "Sessão"


def criar_aplicacao_mock() -> AgendaApplication:
    """Cria uma AgendaApplication com dependências simuladas."""

    aplicacao = AgendaApplication()

    aplicacao.logger = Mock()
    aplicacao.agenda_updater = Mock()

    return aplicacao


def test_atualizar_ano_fluxo_completo(
    tmp_path: Path,
) -> None:
    """Valida o fluxo completo de atualização de ano."""

    arquivo_origem = (
        tmp_path
        / "agenda.docx"
    )

    arquivo_destino = (
        tmp_path
        / "agenda_2027.docx"
    )

    criar_documento_teste(
        arquivo_origem
    )

    aplicacao = AgendaApplication()

    quantidade = aplicacao.atualizar_ano(
        arquivo=arquivo_origem,
        ano_origem=2026,
        ano_destino=2027,
        destino=arquivo_destino,
    )

    assert quantidade == 3
    assert arquivo_destino.exists()

    validar_documento(
        arquivo_destino
    )


def test_atualizar_ano_registra_e_repropaga_erro(
    tmp_path: Path,
) -> None:
    """Valida tratamento de erro na atualização do ano."""

    arquivo_origem = (
        tmp_path
        / "agenda.docx"
    )

    arquivo_destino = (
        tmp_path
        / "agenda_2027.docx"
    )

    aplicacao = criar_aplicacao_mock()

    agenda_updater = cast(
        Any,
        aplicacao.agenda_updater,
    )

    logger = cast(
        Any,
        aplicacao.logger,
    )

    erro = RuntimeError(
        "Erro durante atualização do ano."
    )

    agenda_updater.atualizar_ano.side_effect = erro

    with pytest.raises(
        RuntimeError,
        match="Erro durante atualização do ano.",
    ):
        aplicacao.atualizar_ano(
            arquivo=arquivo_origem,
            ano_origem=2026,
            ano_destino=2027,
            destino=arquivo_destino,
        )

    logger.exception.assert_called_once()


def test_atualizar_datas_registra_e_repropaga_erro(
    tmp_path: Path,
) -> None:
    """Valida tratamento de erro na atualização de datas."""

    arquivo_origem = (
        tmp_path
        / "agenda.docx"
    )

    arquivo_destino = (
        tmp_path
        / "agenda_atualizada.docx"
    )

    aplicacao = criar_aplicacao_mock()

    agenda_updater = cast(
        Any,
        aplicacao.agenda_updater,
    )

    logger = cast(
        Any,
        aplicacao.logger,
    )

    erro = RuntimeError(
        "Erro durante atualização de datas."
    )

    agenda_updater.atualizar_agenda.side_effect = erro

    with pytest.raises(
        RuntimeError,
        match="Erro durante atualização de datas.",
    ):
        aplicacao.atualizar_datas(
            arquivo=arquivo_origem,
            substituicoes={
                "05/01/2026": "05/01/2027",
            },
            destino=arquivo_destino,
        )

    logger.exception.assert_called_once()


def test_atualizar_configuracao_registra_e_repropaga_erro(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Valida tratamento de erro usando configuração."""

    aplicacao = criar_aplicacao_mock()

    logger = cast(
        Any,
        aplicacao.logger,
    )

    erro = RuntimeError(
        "Erro na atualização por configuração."
    )

    def atualizar_ano_mock(
        *args: Any,
        **kwargs: Any,
    ) -> int:
        """Simula uma falha durante atualização."""

        raise erro

    monkeypatch.setattr(
        aplicacao,
        "atualizar_ano",
        atualizar_ano_mock,
    )

    config = AgendaConfig(
        arquivo_origem=Path("agenda.docx"),
        arquivo_destino=Path("agenda_2027.docx"),
        ano_origem=2026,
        ano_destino=2027,
    )

    with pytest.raises(
        RuntimeError,
        match="Erro na atualização por configuração.",
    ):
        aplicacao.atualizar(config)

    logger.exception.assert_called_once()