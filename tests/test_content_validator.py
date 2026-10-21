"""Testes de validação de conteúdo antes da publicação (Sprint 5.12)."""

from pathlib import Path
from unittest.mock import Mock

import pytest
from docx import Document

from core.agenda_updater import AgendaUpdater
from services.content_validator import ContentValidator


def _updater() -> AgendaUpdater:
    updater = AgendaUpdater()
    updater.backup_service.criar_backup = Mock()  # type: ignore[method-assign]
    return updater


def test_ano_valida_paragrafos_e_tabelas(tmp_path: Path) -> None:
    origem = tmp_path / "original.docx"
    destino = tmp_path / "atualizada.docx"
    documento = Document()
    documento.add_paragraph("Audiência em 05/01/2026")
    documento.add_table(rows=1, cols=1).cell(0, 0).text = "10/02/2026"
    documento.save(str(origem))
    assert _updater().atualizar_ano(origem, 2026, 2027, destino) == 2
    resultado = Document(str(destino))
    assert "05/01/2027" in resultado.paragraphs[0].text
    assert resultado.tables[0].cell(0, 0).text == "10/02/2027"


def test_ano_pendente_impede_publicacao(tmp_path: Path) -> None:
    origem = tmp_path / "original.docx"
    destino = tmp_path / "atualizada.docx"
    documento = Document()
    documento.add_paragraph("29/02/2024")
    documento.save(str(origem))
    with pytest.raises(ValueError, match="Validação de conteúdo falhou"):
        _updater().atualizar_ano(origem, 2024, 2025, destino)
    assert not destino.exists()


def test_substituicoes_validadas_antes_de_salvar(tmp_path: Path) -> None:
    origem = tmp_path / "original.docx"
    destino = tmp_path / "atualizada.docx"
    documento = Document()
    documento.add_paragraph("05/01/2026")
    documento.save(str(origem))
    assert _updater().atualizar_agenda(
        origem, {"05/01/2026": "05/01/2027"}, destino
    ) == 1
    assert Document(str(destino)).paragraphs[0].text == "05/01/2027"


def test_validacao_detecta_contagem_incorreta() -> None:
    documento = Document()
    documento.add_paragraph("05/01/2026")
    with pytest.raises(ValueError, match="Validação de conteúdo falhou"):
        ContentValidator.validar_substituicoes(
            documento, {"05/01/2026": "05/01/2027"}, 1, 0
        )


def test_ano_sem_datas_e_valido(tmp_path: Path) -> None:
    origem = tmp_path / "original.docx"
    destino = tmp_path / "atualizada.docx"
    documento = Document()
    documento.add_paragraph("Agenda sem datas")
    documento.save(str(origem))
    assert _updater().atualizar_ano(origem, 2026, 2027, destino) == 0
    assert destino.exists()
