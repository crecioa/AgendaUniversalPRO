from docx import Document

from services.text_updater import TextUpdater


def test_substituicao_fragmentada_com_tamanho_diferente() -> None:
    documento = Document()
    paragrafo = documento.add_paragraph()

    inicio = paragrafo.add_run("Audiência em ")
    inicio.bold = True

    data1 = paragrafo.add_run("10/01/")
    data1.italic = True

    data2 = paragrafo.add_run("2026")
    data2.italic = True

    fim = paragrafo.add_run(" às 14h.")
    fim.bold = True

    quantidade = TextUpdater().atualizar_datas(
        documento,
        {"10/01/2026": "10 de janeiro de 2027"},
    )

    assert quantidade == 1
    assert paragrafo.text == "Audiência em 10 de janeiro de 2027 às 14h."
    assert inicio.bold is True
    assert fim.bold is True
    assert fim.text == " às 14h."
