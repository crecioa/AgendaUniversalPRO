from pathlib import Path

import pytest
from docx import Document

from services.docx_integrity_auditor import DocxIntegrityAuditor


def test_auditoria_docx_valido_sprint71(tmp_path: Path) -> None:
    caminho = tmp_path / "agenda_valida.docx"

    documento = Document()
    documento.add_paragraph("Agenda Universal PRO")
    documento.save(str(caminho))

    relatorio = DocxIntegrityAuditor.auditar(caminho)

    assert relatorio.zip_valido is True
    assert relatorio.total_componentes > 0
    assert "word/document.xml" in relatorio.componentes
    assert relatorio.componentes_xml > 0
    assert relatorio.arquivos_corrompidos == ()


def test_auditoria_docx_invalido_sprint71(tmp_path: Path) -> None:
    caminho = tmp_path / "agenda_invalida.docx"

    caminho.write_bytes(b"Arquivo que nao e um DOCX valido")

    relatorio = DocxIntegrityAuditor.auditar(caminho)

    assert relatorio.zip_valido is False
    assert relatorio.total_componentes == 0


def test_auditoria_docx_inexistente_sprint71(
    tmp_path: Path,
) -> None:
    caminho = tmp_path / "nao_existe.docx"

    with pytest.raises(FileNotFoundError):
        DocxIntegrityAuditor.auditar(caminho)


def test_auditoria_docx_nao_modifica_arquivo_sprint71(
    tmp_path: Path,
) -> None:
    caminho = tmp_path / "agenda_preservada.docx"

    documento = Document()
    documento.add_paragraph("Agenda Universal PRO")
    documento.save(str(caminho))

    conteudo_antes = caminho.read_bytes()

    DocxIntegrityAuditor.auditar(caminho)

    conteudo_depois = caminho.read_bytes()

    assert conteudo_antes == conteudo_depois