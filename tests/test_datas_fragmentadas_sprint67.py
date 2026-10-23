from pathlib import Path

from docx import Document

from services.text_updater import TextUpdater


def test_data_dividida_entre_runs_sprint67(tmp_path: Path) -> None:
    origem = tmp_path / "agenda_fragmentada.docx"

    documento = Document()
    paragrafo = documento.add_paragraph()

    primeiro = paragrafo.add_run("10/01/")
    primeiro.bold = True

    segundo = paragrafo.add_run("2026")
    segundo.italic = True

    documento.save(str(origem))

    documento = Document(str(origem))

    quantidade = TextUpdater().atualizar_datas(
        documento,
        {"10/01/2026": "10/01/2027"},
    )

    paragrafo_atualizado = documento.paragraphs[0]

    assert quantidade == 1
    assert paragrafo_atualizado.text == "10/01/2027"
    assert paragrafo_atualizado.runs[0].bold is True
    assert paragrafo_atualizado.runs[1].italic is True
