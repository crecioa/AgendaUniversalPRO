"""
Agenda Universal PRO
Arquivo: tests/test_agenda_updater_ano.py

Teste da atualização automática de ano através do AgendaUpdater.
"""

from pathlib import Path

from docx import Document

from core.agenda_updater import AgendaUpdater


def criar_documento_teste(arquivo: Path) -> None:
    """Cria um documento Word temporário para o teste."""

    documento = Document()

    documento.add_paragraph("Agenda Universal PRO")
    documento.add_paragraph("Audiência em 05/01/2026.")
    documento.add_paragraph("Reunião em 10/02/2026.")
    documento.add_paragraph("Prazo processual em 15/03/2027.")
    documento.add_paragraph("Este parágrafo não possui nenhuma data.")

    documento.save(str(arquivo))


def main() -> None:
    """Executa o teste de integração."""

    arquivo = Path("modelos") / "teste_agenda_ano.docx"
    destino = Path("modelos") / "teste_agenda_ano_atualizada.docx"

    criar_documento_teste(arquivo)

    agenda_updater = AgendaUpdater()

    quantidade = agenda_updater.atualizar_ano(
        arquivo,
        2026,
        2027,
        destino,
    )

    documento = Document(str(destino))

    textos = [paragrafo.text for paragrafo in documento.paragraphs]

    print("\n=== TESTE DO AGENDAUPDATER - ATUALIZAÇÃO DE ANO ===\n")

    for texto in textos:
        print(texto)

    print(f"\nSubstituições realizadas: {quantidade}")

    assert quantidade == 2

    assert "Audiência em 05/01/2027." in textos
    assert "Reunião em 10/02/2027." in textos
    assert "Prazo processual em 15/03/2027." in textos

    print("\nDocumento atualizado criado com sucesso.")
    print("Teste concluído com sucesso.")


if __name__ == "__main__":
    main()