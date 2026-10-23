
"""
Agenda Universal PRO
Arquivo: services/text_updater.py

Serviço responsável por localizar e atualizar datas e dias da semana
nos parágrafos e tabelas de um documento Word.
"""

from __future__ import annotations

import re
from logging import Logger
from typing import Callable

from docx.document import Document as DocxDocument
from docx.table import _Cell

from core.date_parser import DateParser
from services.logger_service import LoggerService


class TextUpdater:
    """Localiza e atualiza datas em documentos Word."""

    DIAS_SEMANA_PATTERN = (
        r"Segunda-feira|Terça-feira|Quarta-feira|"
        r"Quinta-feira|Sexta-feira|Sábado|Domingo"
    )

    DATA_COM_DIA_PATTERN = re.compile(
        rf"(?P<dia>{DIAS_SEMANA_PATTERN})"
        r"(?P<separador>\s*[,–—-]\s*)"
        r"(?P<data>\d{1,2}/\d{1,2}/\d{4})",
        re.IGNORECASE,
    )

    def __init__(self) -> None:
        self.logger: Logger = LoggerService.get_logger()

    def localizar_datas(
        self,
        documento: DocxDocument,
    ) -> list[tuple[int, str, list[str]]]:
        """
        Percorre os parágrafos procurando datas.

        Returns:
            Lista contendo:
                (índice do parágrafo, texto, datas encontradas)
        """

        ocorrencias: list[tuple[int, str, list[str]]] = []

        self.logger.info(
            "Iniciando leitura dos parágrafos."
        )

        for indice, paragrafo in enumerate(
            documento.paragraphs
        ):
            texto = paragrafo.text.strip()

            if not texto:
                continue

            datas = DateParser.encontrar_datas(texto)

            if datas:
                ocorrencias.append(
                    (indice, texto, datas)
                )

                self.logger.info(
                    "Parágrafo %s: %d data(s) encontrada(s): %s",
                    indice,
                    len(datas),
                    ", ".join(datas),
                )

        self.logger.info(
            "Leitura concluída. %d parágrafo(s) com datas.",
            len(ocorrencias),
        )

        return ocorrencias

    @staticmethod
    def _substituir_texto(
        texto: str,
        substituicoes: dict[str, str],
    ) -> tuple[str, int]:
        """
        Substitui datas em um texto.

        Returns:
            Tupla contendo:
                (texto atualizado, quantidade de substituições)
        """

        texto_atualizado = texto
        quantidade = 0

        for data_antiga, data_nova in substituicoes.items():
            ocorrencias = texto_atualizado.count(
                data_antiga
            )

            if ocorrencias:
                texto_atualizado = texto_atualizado.replace(
                    data_antiga,
                    data_nova,
                )

                quantidade += ocorrencias

        return texto_atualizado, quantidade

    def _atualizar_paragrafo(
        self,
        paragrafo,
        substituicoes: dict[str, str],
    ) -> int:
        """Atualiza datas preservando os runs e sua formata??o."""

        texto_original = paragrafo.text

        if not texto_original.strip():
            return 0

        runs = list(paragrafo.runs)

        if not runs:
            return 0

        # Trata separadamente os runs convencionais e os hyperlinks.
        # N?o mistura caracteres de dentro e fora de um hyperlink.
        if "".join(run.text for run in runs) != texto_original:
            class GrupoRuns:
                def __init__(self, runs_grupo):
                    self.runs = runs_grupo

                @property
                def text(self):
                    return "".join(run.text for run in self.runs)

            quantidade_total = 0

            for elemento in paragrafo.iter_inner_content():
                if hasattr(elemento, "runs"):
                    runs_grupo = list(elemento.runs)
                else:
                    runs_grupo = [elemento]

                if not runs_grupo:
                    continue

                grupo = GrupoRuns(runs_grupo)

                quantidade_total += self._atualizar_paragrafo(
                    grupo,
                    substituicoes,
                )

            return quantidade_total

        texto_atualizado = texto_original
        quantidade = 0

        # Executa as substitui??es em sequ?ncia, mantendo o
        # comportamento original de _substituir_texto.
        for antigo, novo in substituicoes.items():
            if not antigo:
                continue

            posicao = 0

            while True:
                indice = texto_atualizado.find(antigo, posicao)

                if indice == -1:
                    break

                fim = indice + len(antigo)

                # Mapeia as posi??es dos caracteres aos runs atuais.
                limites = []
                cursor = 0

                for run in runs:
                    inicio_run = cursor
                    cursor += len(run.text)
                    limites.append((inicio_run, cursor))

                run_inicial = next(
                    i for i, (_, limite) in enumerate(limites)
                    if limite > indice
                )

                run_final = next(
                    i for i, (_, limite) in enumerate(limites)
                    if limite >= fim
                )

                inicio_primeiro = limites[run_inicial][0]
                inicio_ultimo = limites[run_final][0]

                if run_inicial == run_final:
                    run = runs[run_inicial]
                    deslocamento = indice - inicio_primeiro
                    run.text = (
                        run.text[:deslocamento]
                        + novo
                        + run.text[deslocamento + len(antigo):]
                    )
                else:
                    primeiro = runs[run_inicial]
                    ultimo = runs[run_final]

                    prefixo = primeiro.text[:indice - inicio_primeiro]
                    sufixo = ultimo.text[fim - inicio_ultimo:]

                    # Mant?m a distribui??o dos caracteres substitu?dos
                    # entre os runs sempre que os comprimentos coincidirem.
                    if len(antigo) == len(novo):
                        restante = novo

                        for i in range(run_inicial, run_final + 1):
                            run = runs[i]
                            inicio_atual, fim_atual = limites[i]

                            inicio_trecho = max(indice, inicio_atual)
                            fim_trecho = min(fim, fim_atual)
                            tamanho = fim_trecho - inicio_trecho

                            parte = restante[:tamanho]
                            restante = restante[tamanho:]

                            if i == run_inicial:
                                run.text = prefixo + parte
                            elif i == run_final:
                                run.text = parte + sufixo
                            else:
                                run.text = parte
                    else:
                        primeiro.text = prefixo + novo

                        for i in range(run_inicial + 1, run_final):
                            runs[i].text = ""

                        ultimo.text = sufixo

                texto_atualizado = (
                    texto_atualizado[:indice]
                    + novo
                    + texto_atualizado[fim:]
                )

                quantidade += 1
                posicao = indice + len(novo)

        return quantidade

    def _atualizar_celula(
        self,
        celula: _Cell,
        substituicoes: dict[str, str],
    ) -> int:
        """
        Atualiza as substituições nos parágrafos de uma célula.
        """

        quantidade = 0

        for paragrafo in celula.paragraphs:
            quantidade += self._atualizar_paragrafo(
                paragrafo,
                substituicoes,
            )

        return quantidade

    def atualizar_datas(
        self,
        documento: DocxDocument,
        substituicoes: dict[str, str],
    ) -> int:
        """
        Substitui datas nos parágrafos e tabelas do documento.

        Returns:
            Quantidade de substituições realizadas.
        """

        quantidade = 0

        self.logger.info(
            "Iniciando atualização das datas."
        )

        for paragrafo in documento.paragraphs:
            quantidade += self._atualizar_paragrafo(
                paragrafo,
                substituicoes,
            )

        for tabela in documento.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    quantidade += self._atualizar_celula(
                        celula,
                        substituicoes,
                    )

        # Atualiza cabe?alhos e rodap?s de todas as se??es.
        # Evita processar novamente partes compartilhadas.
        partes_processadas = set()

        for secao in documento.sections:
            for parte in (
                secao.header,
                secao.footer,
                secao.first_page_header,
                secao.first_page_footer,
                secao.even_page_header,
                secao.even_page_footer,
            ):
                identificador = str(parte.part.partname)

                if identificador in partes_processadas:
                    continue

                partes_processadas.add(identificador)

                for paragrafo in parte.paragraphs:
                    quantidade += self._atualizar_paragrafo(
                        paragrafo,
                        substituicoes,
                    )

                for tabela in parte.tables:
                    for linha in tabela.rows:
                        for celula in linha.cells:
                            quantidade += self._atualizar_celula(
                                celula,
                                substituicoes,
                            )

        self.logger.info(
            "Atualização concluída. "
            "%d substituição(ões) realizada(s).",
            quantidade,
        )

        return quantidade

    @staticmethod
    def _atualizar_dia_e_data(
        texto: str,
        ano_origem: int,
        ano_destino: int,
    ) -> tuple[str, int, set[str]]:
        """
        Atualiza pares de dia da semana e data.

        Returns:
            Tupla contendo:
                (texto atualizado, quantidade, datas processadas)
        """

        quantidade = 0
        datas_processadas: set[str] = set()

        def substituir_par(
            match: re.Match[str],
        ) -> str:
            nonlocal quantidade

            data_antiga = match.group("data")
            data_convertida = DateParser.converter_data(
                data_antiga
            )

            if data_convertida is None:
                return match.group(0)

            if data_convertida.year != ano_origem:
                return match.group(0)

            nova_data = DateParser.atualizar_ano(
                data_antiga,
                ano_destino,
            )

            if nova_data is None:
                return match.group(0)

            novo_dia = DateParser.obter_dia_semana(
                nova_data
            )

            if novo_dia is None:
                return match.group(0)

            quantidade += 1
            datas_processadas.add(data_antiga)

            return (
                f"{novo_dia}"
                f"{match.group('separador')}"
                f"{nova_data}"
            )

        texto_atualizado = TextUpdater.DATA_COM_DIA_PATTERN.sub(
            substituir_par,
            texto,
        )

        return (
            texto_atualizado,
            quantidade,
            datas_processadas,
        )

    @staticmethod
    def _atualizar_datas_restantes(
        texto: str,
        ano_origem: int,
        ano_destino: int,
        datas_processadas: set[str],
    ) -> tuple[str, int]:
        """
        Atualiza datas que não possuem o dia da semana associado.
        """

        quantidade = 0
        texto_atualizado = texto

        datas = DateParser.encontrar_datas(texto)

        for data in datas:
            if data in datas_processadas:
                continue

            data_convertida = DateParser.converter_data(data)

            if data_convertida is None:
                continue

            if data_convertida.year != ano_origem:
                continue

            nova_data = DateParser.atualizar_ano(
                data,
                ano_destino,
            )

            if nova_data is None:
                continue

            ocorrencias = texto_atualizado.count(data)

            if ocorrencias:
                texto_atualizado = texto_atualizado.replace(
                    data,
                    nova_data,
                )

                quantidade += ocorrencias

        return texto_atualizado, quantidade

    def atualizar_ano(
        self,
        documento: DocxDocument,
        ano_origem: int,
        ano_destino: int,
    ) -> int:
        """
        Atualiza automaticamente o ano das datas existentes
        nos parágrafos e tabelas do documento.

        Também recalcula o dia da semana quando ele aparece
        associado à data.

        Returns:
            Quantidade de datas atualizadas.
        """

        quantidade = 0

        self.logger.info(
            "Iniciando atualização automática de %d para %d.",
            ano_origem,
            ano_destino,
        )

        def atualizar_texto(
            texto: str,
        ) -> tuple[str, int]:
            texto_atualizado, quantidade_dias, datas_processadas = (
                self._atualizar_dia_e_data(
                    texto,
                    ano_origem,
                    ano_destino,
                )
            )

            texto_atualizado, quantidade_restante = (
                self._atualizar_datas_restantes(
                    texto_atualizado,
                    ano_origem,
                    ano_destino,
                    datas_processadas,
                )
            )

            return (
                texto_atualizado,
                quantidade_dias + quantidade_restante,
            )

        def atualizar_paragrafo(paragrafo) -> int:
            texto_original = paragrafo.text

            if not texto_original.strip():
                return 0

            texto_atualizado, quantidade_local = (
                atualizar_texto(texto_original)
            )

            if texto_atualizado != texto_original:
                runs = list(paragrafo.runs)

                if "".join(run.text for run in runs) != texto_original:
                    # Processa trechos convencionais e hyperlinks
                    # separadamente, preservando seus relacionamentos.
                    quantidade_segura = 0

                    for elemento in paragrafo.iter_inner_content():
                        if hasattr(elemento, "runs"):
                            runs_elemento = list(elemento.runs)
                        else:
                            runs_elemento = [elemento]

                        texto_grupo = "".join(
                            run.text for run in runs_elemento
                        )

                        novo_grupo, qtd = atualizar_texto(
                            texto_grupo
                        )

                        if novo_grupo != texto_grupo:
                            from difflib import SequenceMatcher

                            comparador_grupo = SequenceMatcher(
                                None,
                                texto_grupo,
                                novo_grupo,
                                autojunk=False,
                            )

                            operacoes_grupo = [
                                (i1, i2, novo_grupo[j1:j2])
                                for tipo, i1, i2, j1, j2
                                in comparador_grupo.get_opcodes()
                                if tipo != "equal"
                            ]

                            for i1, i2, novo_trecho in reversed(
                                operacoes_grupo
                            ):
                                limites_grupo = []
                                posicao_grupo = 0

                                for run in runs_elemento:
                                    inicio = posicao_grupo
                                    posicao_grupo += len(run.text)
                                    limites_grupo.append(
                                        (inicio, posicao_grupo)
                                    )

                                afetados = [
                                    indice
                                    for indice, (inicio, fim)
                                    in enumerate(limites_grupo)
                                    if (
                                        inicio < i2 and fim > i1
                                    ) or (
                                        i1 == i2
                                        and inicio <= i1 < fim
                                    )
                                ]

                                if not afetados:
                                    raise ValueError(
                                        "N?o foi poss?vel mapear "
                                        "a altera??o no hyperlink."
                                    )

                                primeiro_indice = afetados[0]
                                ultimo_indice = afetados[-1]

                                primeiro = runs_elemento[
                                    primeiro_indice
                                ]
                                ultimo = runs_elemento[
                                    ultimo_indice
                                ]

                                inicio_primeiro = limites_grupo[
                                    primeiro_indice
                                ][0]
                                inicio_ultimo = limites_grupo[
                                    ultimo_indice
                                ][0]

                                prefixo = primeiro.text[
                                    :i1 - inicio_primeiro
                                ]
                                sufixo = ultimo.text[
                                    i2 - inicio_ultimo:
                                ]

                                if primeiro_indice == ultimo_indice:
                                    primeiro.text = (
                                        prefixo
                                        + novo_trecho
                                        + sufixo
                                    )
                                else:
                                    primeiro.text = (
                                        prefixo + novo_trecho
                                    )

                                    for indice in range(
                                        primeiro_indice + 1,
                                        ultimo_indice,
                                    ):
                                        runs_elemento[indice].text = ""

                                    ultimo.text = sufixo

                            texto_final = "".join(
                                run.text for run in runs_elemento
                            )

                            if texto_final != novo_grupo:
                                raise ValueError(
                                    "Falha ao preservar o texto "
                                    "do hyperlink."
                                )

                        quantidade_segura += qtd

                    return quantidade_segura

                # Distribui altera??es posicionais preservando os runs.
                # Para diferen?as de comprimento, utiliza o primeiro
                # run afetado e mant?m o texto n?o alterado.
                from difflib import SequenceMatcher

                comparador = SequenceMatcher(
                    None,
                    texto_original,
                    texto_atualizado,
                    autojunk=False,
                )

                operacoes = [
                    (i1, i2, texto_atualizado[j1:j2])
                    for tipo, i1, i2, j1, j2
                    in comparador.get_opcodes()
                    if tipo != "equal"
                ]

                for i1, i2, novo_texto in reversed(operacoes):
                    posicao = 0
                    limites = []

                    for run in runs:
                        limites.append((posicao, posicao + len(run.text)))
                        posicao += len(run.text)

                    afetados = [
                        indice
                        for indice, (inicio_run, fim_run)
                        in enumerate(limites)
                        if (
                            inicio_run < i2 and fim_run > i1
                        ) or (
                            i1 == i2 and inicio_run <= i1 < fim_run
                        )
                    ]

                    if not afetados:
                        continue

                    primeiro_indice = afetados[0]
                    ultimo_indice = afetados[-1]

                    primeiro = runs[primeiro_indice]
                    ultimo = runs[ultimo_indice]

                    inicio_primeiro = limites[primeiro_indice][0]
                    inicio_ultimo = limites[ultimo_indice][0]

                    prefixo = primeiro.text[:i1 - inicio_primeiro]
                    sufixo = ultimo.text[i2 - inicio_ultimo:]

                    if primeiro_indice == ultimo_indice:
                        primeiro.text = prefixo + novo_texto + sufixo
                    else:
                        primeiro.text = prefixo + novo_texto

                        for indice in range(
                            primeiro_indice + 1,
                            ultimo_indice,
                        ):
                            runs[indice].text = ""

                        ultimo.text = sufixo

                if paragrafo.text != texto_atualizado:
                    raise ValueError(
                        "Falha ao preservar o texto durante "
                        "a atualiza??o do ano."
                    )

            return quantidade_local

        for paragrafo in documento.paragraphs:
            quantidade += atualizar_paragrafo(
                paragrafo
            )

        for tabela in documento.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    for paragrafo in celula.paragraphs:
                        quantidade += atualizar_paragrafo(
                            paragrafo
                        )

        self.logger.info(
            "Atualização automática concluída. "
            "%d substituição(ões) realizada(s).",
            quantidade,
        )

        # Processa cabecalhos e rodapes sem repetir partes compartilhadas.
        partes_processadas = set()

        for secao in documento.sections:
            for parte in (
                secao.header,
                secao.footer,
                secao.first_page_header,
                secao.first_page_footer,
                secao.even_page_header,
                secao.even_page_footer,
            ):
                identificador = str(parte.part.partname)

                if identificador in partes_processadas:
                    continue

                partes_processadas.add(identificador)

                for paragrafo in parte.paragraphs:
                    quantidade += atualizar_paragrafo(paragrafo)

                for tabela in parte.tables:
                    for linha in tabela.rows:
                        for celula in linha.cells:
                            for paragrafo in celula.paragraphs:
                                quantidade += atualizar_paragrafo(
                                    paragrafo
                                )

        return quantidade