"""Relatório detalhado e exportação opcional de atualizações da agenda."""

from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path

from docx.document import Document as DocxDocument


@dataclass(frozen=True)
class Alteracao:
    """Texto de um local do documento antes e depois do processamento."""

    local: str
    antes: str
    depois: str
    datas_antes: list[str]
    datas_depois: list[str]


@dataclass(frozen=True)
class UpdateReport:
    origem: str
    destino: str
    modo: str
    previstas: int
    atualizadas: int
    pendentes: int
    status: str
    alteracoes: list[Alteracao]


class UpdateReportService:
    """Gera relatórios para os locais atendidos pelo TextUpdater."""

    DATA = re.compile(r"(?<!\d)\d{1,2}/\d{1,2}/\d{4}(?!\d)")

    @staticmethod
    def capturar(documento: DocxDocument) -> dict[str, str]:
        """Captura parágrafos e células das tabelas de nível superior."""
        textos: dict[str, str] = {}
        for indice, paragrafo in enumerate(documento.paragraphs):
            textos[f"paragrafo[{indice}]"] = paragrafo.text
        for ti, tabela in enumerate(documento.tables):
            for li, linha in enumerate(tabela.rows):
                for ci, celula in enumerate(linha.cells):
                    for pi, paragrafo in enumerate(celula.paragraphs):
                        textos[f"tabela[{ti}]/linha[{li}]/celula[{ci}]/paragrafo[{pi}]"] = paragrafo.text
        return textos

    @classmethod
    def criar(
        cls,
        origem: str | Path,
        destino: str | Path,
        modo: str,
        antes: dict[str, str],
        depois: dict[str, str],
        previstas: int | None,
        atualizadas: int,
    ) -> UpdateReport:
        """Produz resumo e detalhes; só deve ser chamado após validação."""
        alteracoes = [
            Alteracao(
                local=local,
                antes=texto,
                depois=depois.get(local, ""),
                datas_antes=cls.DATA.findall(texto),
                datas_depois=cls.DATA.findall(depois.get(local, "")),
            )
            for local, texto in antes.items()
            if texto != depois.get(local, "")
        ]
        total = atualizadas if previstas is None else previstas
        pendentes = max(0, total - atualizadas)
        return UpdateReport(
            origem=str(origem), destino=str(destino), modo=modo,
            previstas=total, atualizadas=atualizadas, pendentes=pendentes,
            status="sucesso" if pendentes == 0 else "pendente",
            alteracoes=alteracoes,
        )

    @staticmethod
    def exportar(relatorio: UpdateReport, destino: str | Path) -> Path:
        """Grava JSON UTF-8 com criação exclusiva, sem sobrescrever."""
        caminho = Path(destino)
        if caminho.suffix.lower() != ".json":
            raise ValueError("O relatório deve possuir extensão .json.")
        conteudo = json.dumps(asdict(relatorio), ensure_ascii=False, indent=2)
        with caminho.open("x", encoding="utf-8") as arquivo:
            arquivo.write(conteudo + "\n")
        return caminho
