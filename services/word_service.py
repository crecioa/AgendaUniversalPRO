"""
Agenda Universal PRO
Arquivo: services/word_service.py

Serviço responsável pela manipulação de documentos Word (.docx).
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument

from services.logger_service import LoggerService


class WordService:
    """Serviço para abrir e salvar documentos Word."""

    def __init__(self) -> None:
        self.logger = LoggerService.get_logger()

    def abrir_documento(self, arquivo: str | Path) -> DocxDocument:
        """Abre um documento Word."""
        caminho = Path(arquivo)

        if not caminho.exists():
            raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")

        self.logger.info("Abrindo documento: %s", caminho.name)
        return Document(str(caminho))

    def salvar_documento(
        self,
        documento: DocxDocument,
        destino: str | Path,
    ) -> None:
        """Salva em arquivo temporário e publica sem sobrescrever destino."""
        caminho = Path(destino)

        if caminho.exists():
            raise FileExistsError(f"O arquivo de destino já existe: {caminho}")

        descritor, nome_temporario = tempfile.mkstemp(
            prefix=f".{caminho.stem}.",
            suffix=".docx",
            dir=str(caminho.parent),
        )
        os.close(descritor)
        temporario = Path(nome_temporario)

        try:
            documento.save(str(temporario))
            # Um hard link cria o destino somente se ele ainda não existir.
            # Diferentemente de os.replace(), não sobrescreve arquivos.
            os.link(temporario, caminho)
        finally:
            temporario.unlink(missing_ok=True)

        self.logger.info("Documento salvo: %s", caminho.name)

    @staticmethod
    def contar_paragrafos(documento: DocxDocument) -> int:
        """Retorna a quantidade de parágrafos."""
        return len(documento.paragraphs)

    @staticmethod
    def contar_tabelas(documento: DocxDocument) -> int:
        """Retorna a quantidade de tabelas."""
        return len(documento.tables)
