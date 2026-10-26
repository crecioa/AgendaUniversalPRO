"""
Agenda Universal PRO
Sprint 7.0 — Auditoria estrutural de documentos Word.
"""

from __future__ import annotations

from dataclasses import dataclass

from docx.document import Document as DocxDocument
from docx.oxml.ns import qn


@dataclass(frozen=True)
class RelatorioAuditoria:
    """Resumo da estrutura de um documento Word."""

    paragrafos: int
    tabelas: int
    secoes: int
    hyperlinks: int
    imagens: int
    campos_especiais: int
    cabecalhos_com_conteudo: int
    rodapes_com_conteudo: int


class DocumentAuditor:
    """Analisa a estrutura de documentos Word sem modificá-los."""

    @staticmethod
    def _contar_elementos(documento: DocxDocument, tag: str) -> int:
        """Conta elementos XML no corpo principal do documento."""
        return sum(
            1
            for elemento in documento.element.body.iter()
            if elemento.tag == qn(tag)
        )

    @staticmethod
    def _possui_conteudo(parte) -> bool:
        """Verifica se uma parte possui texto, tabelas ou imagens."""
        if any(paragrafo.text.strip() for paragrafo in parte.paragraphs):
            return True

        if parte.tables:
            return True

        return any(
            elemento.tag == qn("w:drawing")
            for elemento in parte._element.iter()
        )

    def auditar(self, documento: DocxDocument) -> RelatorioAuditoria:
        """Retorna um relatório estrutural do documento."""

        cabecalhos = sum(
            1
            for secao in documento.sections
            if not secao.header.is_linked_to_previous
            and self._possui_conteudo(secao.header)
        )

        rodapes = sum(
            1
            for secao in documento.sections
            if not secao.footer.is_linked_to_previous
            and self._possui_conteudo(secao.footer)
        )

        return RelatorioAuditoria(
            paragrafos=len(documento.paragraphs),
            tabelas=len(documento.tables),
            secoes=len(documento.sections),
            hyperlinks=self._contar_elementos(
                documento, "w:hyperlink"
            ),
            imagens=self._contar_elementos(
                documento, "w:drawing"
            ),
            campos_especiais=(
                self._contar_elementos(documento, "w:fldChar")
                + self._contar_elementos(documento, "w:fldSimple")
            ),
            cabecalhos_com_conteudo=cabecalhos,
            rodapes_com_conteudo=rodapes,
        )