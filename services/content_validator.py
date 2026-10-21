"""Conferência independente de datas antes da publicação da agenda."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime

from docx.document import Document as DocxDocument


@dataclass(frozen=True)
class RelatorioValidacao:
    """Resumo da conferência das datas do documento."""

    previstas: int
    atualizadas: int
    pendentes: int


class ContentValidator:
    """Confere parágrafos e células das tabelas de nível superior."""

    DATA = re.compile(r"(?<!\d)(\d{1,2}/\d{1,2}/\d{4})(?!\d)")

    @staticmethod
    def _textos(documento: DocxDocument) -> list[str]:
        textos = [paragrafo.text for paragrafo in documento.paragraphs]
        for tabela in documento.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    textos.extend(paragrafo.text for paragrafo in celula.paragraphs)
        return textos

    @classmethod
    def contar_substituicoes(
        cls, documento: DocxDocument, substituicoes: dict[str, str]
    ) -> int:
        return sum(
            texto.count(antiga)
            for texto in cls._textos(documento)
            for antiga in substituicoes
            if antiga
        )

    @classmethod
    def validar_substituicoes(
        cls,
        documento: DocxDocument,
        substituicoes: dict[str, str],
        previstas: int,
        atualizadas: int,
    ) -> RelatorioValidacao:
        # Uma substituição que conserva o próprio texto não é verificável.
        pendentes = sum(
            texto.count(antiga)
            for texto in cls._textos(documento)
            for antiga, nova in substituicoes.items()
            if antiga and antiga != nova and antiga not in nova
        )
        relatorio = RelatorioValidacao(previstas, atualizadas, pendentes)
        if atualizadas != previstas or pendentes:
            raise ValueError(
                "Validação de conteúdo falhou: "
                f"previstas={previstas}, atualizadas={atualizadas}, "
                f"pendentes={pendentes}."
            )
        return relatorio

    @classmethod
    def contar_ano(cls, documento: DocxDocument, ano: int) -> int:
        quantidade = 0
        for texto in cls._textos(documento):
            for match in cls.DATA.finditer(texto):
                data = match.group(1)
                try:
                    convertida = datetime.strptime(data, "%d/%m/%Y")
                except ValueError:
                    continue
                if convertida.year == ano:
                    quantidade += 1
        return quantidade

    @classmethod
    def validar_ano(
        cls,
        documento: DocxDocument,
        ano_origem: int,
        previstas: int,
        atualizadas: int,
    ) -> RelatorioValidacao:
        pendentes = cls.contar_ano(documento, ano_origem)
        relatorio = RelatorioValidacao(previstas, atualizadas, pendentes)
        if atualizadas != previstas or pendentes:
            raise ValueError(
                "Validação de conteúdo falhou: "
                f"previstas={previstas}, atualizadas={atualizadas}, "
                f"pendentes={pendentes}."
            )
        return relatorio
